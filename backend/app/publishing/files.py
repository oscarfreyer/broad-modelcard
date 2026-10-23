"""Files for maintainers: upload templates (blank or with the current data) and
annotated copies of an upload with its problems marked.

* CSV files start with ``#`` comment lines, which uploads skip.
* XLSX templates have an Instructions sheet, the Data sheet (with dropdowns for
  allowed values), and a hidden Lists sheet that feeds the dropdowns.
* Annotated files add a Problems column; XLSX ones also highlight each cell with
  a problem (red: blocking error, yellow: note) and attach a comment. The
  Problems column is ignored on upload, so a fixed annotated file can be
  uploaded again as it is.
* Text that a spreadsheet program could run as a formula (it starts with ``=``,
  ``+``, ``-``, ``@``, a tab, or a carriage return) is written with a leading
  apostrophe, and XLSX cells holding text are always string cells
  (``escape_formula``, CONTRACT.md §11). Numbers stay numbers; uploads drop the
  apostrophe again.
"""

from __future__ import annotations

import csv
import re
from pathlib import Path
from typing import Any, Iterable, Iterator

from ..importers.spreadsheets import escape_formula, open_table
from ..schema.contract import GRADES, NO_INTENDED_USE_LEVEL, SYNTHETIC_DATA_NOTICE, TOOL_SEPARATOR
from ..schema.templates import (
    ACTION_LABELS,
    AI_TYPE_LABELS,
    CHAPTER_PREFIX,
    CODE_SEPARATOR,
    HARM_LABELS,
    INCIDENT_STATUS_LABELS,
    INCIDENT_TYPE_LABELS,
    SOURCE_TYPE_LABELS,
    GUIDELINE_FOLLOWED_LABELS,
    REPORTING_GUIDELINE_LABELS,
    RISK_OF_BIAS_LABELS,
    SCOPE_LABELS,
    STUDY_CONDUCTOR_LABELS,
    STUDY_DESIGN_LABELS,
    APPRAISAL_TOOL_LABELS,
    APPLICABILITY_LABELS,
    QUALITY_TOOL_LABELS,
    SETTING_CATEGORY_LABELS,
    SETTING_LABELS,
    EFFECT_DIRECTION_LABELS,
    EFFECT_JUDGEMENT_LABELS,
    EFFECT_MEASURE_LABELS,
    OUTCOME_TYPE_LABELS,
    BUILT_ON_SEPARATOR,
    CONDITION_SEPARATOR,
    DATASET_TITLES,
    ENTRY_KIND_LABELS,
    FUNCTION_CATEGORY_LABELS,
    INSTRUCTIONS,
    INTENDED_USE_LEVEL_LABELS,
    MODEL_STATUSES,
    PATIENT_DATA_NOTICE,
    PREVALENCE_TIER_LABELS,
    PROBLEMS_COLUMN,
    ROW_DESCRIPTIONS,
    TEMPLATES,
    VALIDATION_TYPE_LABELS,
    YES,
    YES_NO,
    NO,
)
from ..services.atlas import Atlas
from ..validation.issues import Mark
from . import content as contents

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
CSV_MEDIA_TYPE = "text/csv; charset=utf-8"

ERROR_FILL = "FFC7CE"
NOTE_FILL = "FFEB9C"
HEADER_FILL = "E8EEF7"
AUTHOR = "Broad Modelcard"


def choice_lists(atlas: Atlas) -> dict[str, tuple[str, ...]]:
    axes = atlas.axes
    return {
        "lens:age": axes.age,
        "lens:sex": axes.sex,
        "lens:ethnicity": axes.ethnicity,
        "grades": tuple(GRADES),
        "validation_types": tuple(VALIDATION_TYPE_LABELS.values()),
        "yes_no": YES_NO,
        "intended_use_levels": tuple(INTENDED_USE_LEVEL_LABELS.values()),
        "prevalence_tiers": tuple(PREVALENCE_TIER_LABELS.values()),
        "model_statuses": MODEL_STATUSES,
        "entry_kinds": tuple(ENTRY_KIND_LABELS.values()),
        "ai_types": tuple(AI_TYPE_LABELS.values()),
        "function_categories": tuple(FUNCTION_CATEGORY_LABELS.values()),
        "actions": tuple(ACTION_LABELS.values()),
        "incident_types": tuple(INCIDENT_TYPE_LABELS.values()),
        "harms": tuple(HARM_LABELS.values()),
        "incident_statuses": tuple(INCIDENT_STATUS_LABELS.values()),
        "source_types": tuple(SOURCE_TYPE_LABELS.values()),
        "study_designs": tuple(STUDY_DESIGN_LABELS.values()),
        "study_conductors": tuple(STUDY_CONDUCTOR_LABELS.values()),
        "reporting_guidelines": tuple(REPORTING_GUIDELINE_LABELS.values()),
        "guideline_followed": tuple(GUIDELINE_FOLLOWED_LABELS.values()),
        "risk_of_bias": tuple(RISK_OF_BIAS_LABELS.values()),
        "scopes": tuple(SCOPE_LABELS.values()),
        "applicability": tuple(APPLICABILITY_LABELS.values()),
        "quality_tools": tuple(QUALITY_TOOL_LABELS.values()),
        "setting_categories": tuple(SETTING_CATEGORY_LABELS.values()),
        "settings": tuple(dict.fromkeys(SETTING_LABELS.values())),
        "outcome_types": tuple(OUTCOME_TYPE_LABELS.values()),
        "effect_measures": tuple(EFFECT_MEASURE_LABELS.values()),
        "effect_directions": tuple(EFFECT_DIRECTION_LABELS.values()),
        "appraisal_tools": tuple(APPRAISAL_TOOL_LABELS.values()),
        "effect_judgements": tuple(EFFECT_JUDGEMENT_LABELS.values()),
    }


# ------------------------------------------------------------- data rows

def _metric_cells(values: tuple) -> list[Any]:
    (perf, ci_low, ci_high, grade, samples, sites, studies, peer, validation_type, in_use, coverage, level, warnings) = values
    return [
        perf,
        ci_low,
        ci_high,
        grade,
        samples,
        sites,
        studies,
        peer,
        VALIDATION_TYPE_LABELS.get(validation_type, validation_type),
        # A row without an intended-use value (v1.11) leaves the three cells empty.
        None if level == NO_INTENDED_USE_LEVEL else YES if in_use else NO,
        coverage,
        None if level == NO_INTENDED_USE_LEVEL else INTENDED_USE_LEVEL_LABELS.get(level, level),
        "; ".join(warnings),
    ]


def _source_cells(sources: tuple | None, key) -> list[Any]:
    """Action (empty) and the row's Source, Study, and Source date (contract v1.7)."""
    if sources is None:
        return [None, None, None, None]
    default, explicit = sources
    source = explicit.get(key, default)
    return [None, source.source, source.study, source.date]


def data_rows(content: contents.Content | None, atlas: Atlas, sources: tuple | None = None) -> Iterator[list[Any]]:
    """The rows of ``content`` in template column order. ``sources`` is a
    metrics or prevalence set's ``(default source, explicit sources)``
    (``content.load_sources``); without it the source columns stay empty."""
    if content is None:
        return
    names = {concept.concept_id: concept.pt for concept in atlas.concepts}
    lens = atlas.axes.lens_for_index
    if content.kind in ("metrics", "local_metrics"):
        for (combo_idx, concept_id), values in sorted(content.rows.items(), key=lambda item: (item[0][1], item[0][0])):
            group = lens(combo_idx)
            yield [
                concept_id, names.get(concept_id), group.age, group.sex, group.ethnicity, SETTING_LABELS[group.setting],
                *_metric_cells(values), *_source_cells(sources, (combo_idx, concept_id)),
            ]
    elif content.kind == "safety":
        for incident in content.incidents:
            chapter = atlas.by_id.get(incident.chapter_id) if incident.chapter_id else None
            codes = f"{CHAPTER_PREFIX} {chapter.chapter_no}" if chapter else f"{CODE_SEPARATOR} ".join(incident.codes)
            yield [
                incident.incident_id,
                incident.model_version,
                incident.use_case,
                codes,
                incident.age,
                incident.sex,
                incident.ethnicity,
                INCIDENT_TYPE_LABELS.get(incident.incident_type, incident.incident_type),
                HARM_LABELS.get(incident.harm, incident.harm),
                INCIDENT_STATUS_LABELS.get(incident.status, incident.status),
                SOURCE_TYPE_LABELS.get(incident.source_type, incident.source_type),
                incident.source_reference,
                incident.study_id,
                incident.reported_date,
                incident.updated_date,
                incident.summary,
                SCOPE_LABELS.get(incident.scope, incident.scope),
                None,
            ]
    elif content.kind == "studies":
        for study in content.studies:
            yield [
                study.study_id,
                study.title,
                STUDY_DESIGN_LABELS.get(study.design, study.design),
                study.setting,
                study.sites,
                study.patients,
                YES if study.peer_reviewed else NO,
                STUDY_CONDUCTOR_LABELS.get(study.conducted_by, study.conducted_by),
                REPORTING_GUIDELINE_LABELS.get(study.reporting_guideline, study.reporting_guideline),
                GUIDELINE_FOLLOWED_LABELS.get(study.guideline_followed, study.guideline_followed),
                RISK_OF_BIAS_LABELS.get(study.risk_of_bias, study.risk_of_bias),
                study.risk_of_bias_notes,
                study.population,
                study.evaluation_conditions,
                study.model_version,
                study.publication_date,
                study.source_url,
                SCOPE_LABELS.get(study.scope, study.scope),
                study.population_id,
                APPLICABILITY_LABELS.get(study.applicability, study.applicability),
                study.applicability_notes,
                QUALITY_TOOL_LABELS.get(study.quality_tool, study.quality_tool),
                study.served_quality_score,
                SETTING_CATEGORY_LABELS.get(study.setting_category, study.setting_category),
                study.configuration_id,
                None,
            ]
    elif content.kind == "effects":
        for effect in content.effects:
            chapter = atlas.by_id.get(effect.chapter_id) if effect.chapter_id else None
            codes = f"{CHAPTER_PREFIX} {chapter.chapter_no}" if chapter else f"{CODE_SEPARATOR} ".join(effect.codes)
            yield [
                effect.effect_id,
                effect.study_id,
                effect.model_version,
                effect.use_case,
                codes,
                effect.outcome,
                OUTCOME_TYPE_LABELS.get(effect.outcome_type, effect.outcome_type),
                effect.population,
                effect.comparator,
                EFFECT_MEASURE_LABELS.get(effect.effect_measure, effect.effect_measure),
                effect.estimate,
                effect.ci_low,
                effect.ci_high,
                effect.unit,
                EFFECT_DIRECTION_LABELS.get(effect.direction, effect.direction),
                effect.follow_up,
                APPRAISAL_TOOL_LABELS.get(effect.appraisal_tool, effect.appraisal_tool),
                EFFECT_JUDGEMENT_LABELS.get(effect.judgement, effect.judgement),
                effect.notes,
                None,
            ]
    elif content.kind == "multimorbidity":
        for combination in content.combinations:
            codes = f" {CONDITION_SEPARATOR} ".join(combination.condition_ids)
            labels = f" {CONDITION_SEPARATOR} ".join(names.get(concept_id, concept_id) for concept_id in combination.condition_ids)
            for combo_idx, values in sorted(combination.rows.items()):
                group = lens(combo_idx)
                yield [codes, combination.label, labels, group.age, group.sex, group.ethnicity, *_metric_cells(values)]
    elif content.kind == "models":
        for model in content.models:
            for use_case in model.use_cases:
                configuration = model.configuration(use_case)
                yield [
                    model.model_version,
                    model.display_name,
                    model.model_name,
                    model.model_release,
                    model.status,
                    use_case,
                    YES if model.default_use_case == use_case else NO,
                    YES if content.default_model_version == model.model_version else NO,
                    ENTRY_KIND_LABELS.get(model.kind, model.kind),
                    AI_TYPE_LABELS.get(model.ai_type, model.ai_type),
                    model.developer,
                    model.release_date,
                    model.description,
                    f"{BUILT_ON_SEPARATOR} ".join(model.built_on) or None,
                    FUNCTION_CATEGORY_LABELS.get(configuration.function_category, configuration.function_category),
                    configuration.declared_intended_purpose,
                    configuration.declared_regulatory_status,
                    configuration.declared_contact,
                    configuration.declared_documentation_url,
                    YES if configuration.intended_use_defined else NO,
                    # Contract v1.14 (§28).
                    model.replaces,
                    configuration.configuration_id,
                    configuration.configuration_date,
                    configuration.prompt_summary,
                    configuration.prompt_sha256,
                    configuration.retrieval,
                    f"{TOOL_SEPARATOR} ".join(configuration.tools) or None,
                    configuration.sampling_temperature,
                    configuration.sampling_top_p,
                    configuration.sampling_max_output_tokens,
                    configuration.fine_tuning,
                    configuration.performance_measure,
                    configuration.reference_standard,
                    configuration.declared_data_use,
                    configuration.declared_data_location,
                    configuration.declared_retention,
                    configuration.declared_privacy_contact,
                ]
    else:
        for code, tier in content.tiers.items():
            yield [code, names.get(code), PREVALENCE_TIER_LABELS.get(tier, tier), *_source_cells(sources, code)]


# --------------------------------------------------------------------- CSV

def _csv_value(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, float):
        return repr(value)
    return escape_formula(value)


def _comment_lines(kind: str, extra: Iterable[str] = ()) -> list[list[str]]:
    lines = [
        f"# Broad Modelcard: {DATASET_TITLES[kind]} template ({ROW_DESCRIPTIONS[kind]}). {PATIENT_DATA_NOTICE}",
        "# Lines starting with # are ignored. Keep the header row below; the XLSX template explains every column.",
        *extra,
    ]
    return [[line] for line in lines]


def write_template_csv(path: Path, kind: str, rows: Iterable[list[Any]]) -> None:
    with open(path, "w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerows(_comment_lines(kind))
        writer.writerow([column.header for column in TEMPLATES[kind]])
        for row in rows:
            writer.writerow([_csv_value(value) for value in row])


# -------------------------------------------------------------------- XLSX

def _styles():
    from openpyxl.styles import Alignment, Font, PatternFill

    return {
        "header_font": Font(bold=True),
        "header_fill": PatternFill("solid", fgColor=HEADER_FILL),
        "title_font": Font(bold=True, size=14),
        "notice_font": Font(bold=True, color="9C0006"),
        "bold": Font(bold=True),
        "wrap": Alignment(wrap_text=True, vertical="top"),
        "error_fill": PatternFill("solid", fgColor=ERROR_FILL),
        "note_fill": PatternFill("solid", fgColor=NOTE_FILL),
    }


def _xlsx_value(sheet, value: Any) -> Any:
    """A value for a write-only sheet: text that could be read as a formula or an
    error code is escaped and forced to a string cell; other values pass as they are."""
    if not isinstance(value, str):
        return value
    escaped = escape_formula(value)
    if escaped is value and not value.startswith("#"):
        return value  # openpyxl stores it as a string cell
    from openpyxl.cell import WriteOnlyCell

    cell = WriteOnlyCell(sheet, value=escaped)
    cell.data_type = "s"
    return cell


def _xlsx_cell(sheet, value: Any):
    """A styled cell: like ``_xlsx_value``, but always a cell object."""
    from openpyxl.cell import WriteOnlyCell

    cell = WriteOnlyCell(sheet, value=escape_formula(value))
    if isinstance(cell.value, str):
        cell.data_type = "s"
    return cell


def _letter(index: int) -> str:
    from openpyxl.utils import get_column_letter

    return get_column_letter(index + 1)


def _instructions_sheet(workbook, kind: str, lists: dict[str, tuple[str, ...]], styles, known_warnings) -> None:
    sheet = workbook.create_sheet("Instructions")
    sheet.column_dimensions["A"].width = 26
    sheet.column_dimensions["B"].width = 11
    sheet.column_dimensions["C"].width = 70
    sheet.column_dimensions["D"].width = 60

    def cell(value, style=None, wrap=False):
        item = _xlsx_cell(sheet, value)
        if style:
            item.font = styles[style]
        if wrap:
            item.alignment = styles["wrap"]
        return item

    sheet.append([cell(f"Broad Modelcard: {DATASET_TITLES[kind]}", "title_font")])
    sheet.append([cell(PATIENT_DATA_NOTICE, "notice_font")])
    sheet.append([cell(SYNTHETIC_DATA_NOTICE)])
    sheet.append([])
    sheet.append([cell("What this file holds", "bold")])
    sheet.append([cell(f"Fill in the Data sheet with {ROW_DESCRIPTIONS[kind]}.")])
    for line in INSTRUCTIONS[kind]:
        sheet.append([cell(line)])
    sheet.append([])
    sheet.append([cell("How to update", "bold")])
    for step in (
        "1. Fill in or edit the Data sheet. Keep the header row exactly as it is.",
        "2. Save the file as .xlsx (or as CSV UTF-8) and upload it on this data set's page.",
        "3. Fix any problems the check reports (download the marked copy to see them), then upload again.",
        "4. Review the changes, preview them on the atlas, and publish.",
    ):
        sheet.append([cell(step)])
    sheet.append([])
    sheet.append([cell("Columns", "bold")])
    sheet.append([cell(name, "bold") for name in ("Column", "Required", "What to enter", "Allowed values")])
    for column in TEMPLATES[kind]:
        required = "no (ignored)" if column.info else "yes" if column.required else "no (optional column)"
        allowed = ", ".join(lists[column.choices]) if column.choices else ""
        sheet.append([cell(column.header), cell(required), cell(column.help, wrap=True), cell(allowed, wrap=True)])
    if kind in ("metrics", "multimorbidity", "local_metrics"):
        sheet.append([])
        sheet.append([cell("Known warning codes", "bold")])
        sheet.append([cell(", ".join(sorted(known_warnings)), wrap=True)])


def write_template_xlsx(
    path: Path,
    kind: str,
    rows: Iterable[list[Any]],
    atlas: Atlas,
    known_warnings: Iterable[str] = (),
) -> None:
    import openpyxl
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.worksheet.datavalidation import DataValidation

    styles = _styles()
    lists = choice_lists(atlas)
    workbook = openpyxl.Workbook(write_only=True)
    _instructions_sheet(workbook, kind, lists, styles, known_warnings)

    columns = TEMPLATES[kind]
    data = workbook.create_sheet("Data")
    data.freeze_panes = "A2"
    for index, column in enumerate(columns):
        data.column_dimensions[_letter(index)].width = column.width

    list_names = sorted({column.choices for column in columns if column.choices})
    for index, column in enumerate(columns):
        if not column.choices:
            continue
        list_column = _letter(list_names.index(column.choices))
        validation = DataValidation(
            type="list",
            formula1=f"Lists!${list_column}$1:${list_column}${len(lists[column.choices])}",
            allow_blank=True,
            showErrorMessage=True,
            errorStyle="warning",
            errorTitle="Value not in the list",
            error="This value is not in the list. Uploads accept the listed values (and their codes).",
        )
        validation.add(f"{_letter(index)}2:{_letter(index)}1048576")
        data.data_validations.append(validation)

    header = []
    for column in columns:
        item = WriteOnlyCell(data, value=column.header)
        item.font = styles["header_font"]
        item.fill = styles["header_fill"]
        header.append(item)
    data.append(header)
    for row in rows:
        data.append([_xlsx_value(data, value) for value in row])

    list_sheet = workbook.create_sheet("Lists")
    list_sheet.sheet_state = "hidden"
    longest = max((len(lists[name]) for name in list_names), default=0)
    for position in range(longest):
        list_sheet.append([_xlsx_value(list_sheet, lists[name][position]) if position < len(lists[name]) else None for name in list_names])
    workbook.save(path)


# --------------------------------------------------------------- annotated

def _marks_by_row(marks: Iterable[Mark]) -> dict[int, list[Mark]]:
    grouped: dict[int, list[Mark]] = {}
    for mark in marks:
        grouped.setdefault(mark.row, []).append(mark)
    return grouped


def _problem_text(row: int, marks: list[Mark]) -> str:
    parts = [f"{mark.column}: {mark.problem}" if mark.column else mark.problem for mark in marks]
    return f"Row {row}: " + " | ".join(parts)


def _report_lines(report: dict, truncated: bool, header_row: int) -> list[str]:
    """Every problem group, plus the problems of the whole file and the header row."""
    lines = []
    for label, groups in (("Must fix", report.get("errors", [])), ("Note", report.get("notes", []))):
        for group in groups:
            lines.append(f"{label}: {group['message']} ({group['count']:,}). {group['fix']}")
            for example in group["examples"]:
                if example["row"] == 0:
                    lines.append(f"  Whole file: {example['problem']}")
                elif example["row"] == header_row:
                    lines.append(f"  Header row: {example['problem']}")
    if truncated:
        lines.append("Only the first problems are marked in the rows; fix them and check the file again.")
    return lines


def write_annotated_csv(path: Path, source: Path, source_format: str, kind: str, report: dict, marks: list[Mark], truncated: bool) -> None:
    by_row = _marks_by_row(marks)
    with open_table(source, source_format) as table, open(path, "w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        notes = [f"# {line}" for line in _report_lines(report, truncated, table.header_row)]
        writer.writerows(_comment_lines(kind, ["# Checked copy: the Problems column lists what to fix in each row.", *notes]))
        writer.writerow([PROBLEMS_COLUMN.header, *(_csv_value(name) for name in table.header)])
        for number, cells in table.rows():
            row_marks = by_row.get(number)
            writer.writerow([_problem_text(number, row_marks) if row_marks else "", *(_csv_value(value) for value in cells)])


def write_annotated_xlsx(path: Path, source: Path, source_format: str, kind: str, report: dict, marks: list[Mark], truncated: bool) -> None:
    import openpyxl
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.comments import Comment

    styles = _styles()
    by_row = _marks_by_row(marks)
    workbook = openpyxl.Workbook(write_only=True)
    summary = workbook.create_sheet("Summary")
    summary.column_dimensions["A"].width = 120
    title = WriteOnlyCell(summary, value=f"Broad Modelcard: checked copy of a {DATASET_TITLES[kind]} upload")
    title.font = styles["title_font"]
    summary.append([title])
    notice = WriteOnlyCell(summary, value=PATIENT_DATA_NOTICE)
    notice.font = styles["notice_font"]
    summary.append([notice])
    summary.append(["Red cells must be fixed before publishing; yellow cells are notes. Hover a cell to read its comment."])
    summary.append(["The Problems column is ignored when you upload this file again."])
    summary.append([])

    data = workbook.create_sheet("Data")
    data.freeze_panes = "B2"
    data.column_dimensions["A"].width = PROBLEMS_COLUMN.width
    with open_table(source, source_format) as table:
        for line in _report_lines(report, truncated, table.header_row):
            summary.append([_xlsx_value(summary, line)])
        positions = {str(name): index for index, name in enumerate(table.header) if name is not None}
        header = []
        for value in [PROBLEMS_COLUMN.header, *table.header]:
            item = _xlsx_cell(data, value)
            item.font = styles["header_font"]
            item.fill = styles["header_fill"]
            header.append(item)
        for name, index in positions.items():
            matched = next((column for column in TEMPLATES[kind] if column.header == name), None)
            data.column_dimensions[_letter(index + 1)].width = matched.width if matched else 14
        data.append(header)
        for number, cells in table.rows():
            row_marks = by_row.get(number)
            if not row_marks:
                data.append([None, *(_xlsx_value(data, value) for value in cells)])
                continue
            out = [_xlsx_cell(data, value) for value in [_problem_text(number, row_marks), *cells]]
            out[0].fill = styles["error_fill"] if any(mark.blocking for mark in row_marks) else styles["note_fill"]
            by_cell: dict[int, list[Mark]] = {}
            for mark in row_marks:
                if mark.column in positions:
                    by_cell.setdefault(positions[mark.column] + 1, []).append(mark)
            for index, cell_marks in by_cell.items():
                while index >= len(out):
                    out.append(WriteOnlyCell(data, value=None))
                out[index].fill = styles["error_fill"] if any(mark.blocking for mark in cell_marks) else styles["note_fill"]
                comment = Comment("\n".join(mark.problem for mark in cell_marks), AUTHOR)
                comment.width, comment.height = 320, 90
                out[index].comment = comment
            data.append(out)
    workbook.save(path)


def write_annotated(path: Path, file_format: str, **kwargs) -> None:
    (write_annotated_xlsx if file_format == "xlsx" else write_annotated_csv)(path, **kwargs)


def template_filename(kind: str, file_format: str, *, target: tuple[str, ...] | None = None, version: int | None = None) -> str:
    """``target``: a card (model version, use case), or a population id as a 1-tuple (contract v1.10)."""
    parts = ["broad-modelcard", kind]
    if target:
        parts += [_slug(part) for part in target]
    parts.append(f"v{version}" if version is not None else "template")
    return "-".join(parts) + f".{file_format}"


def _slug(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", value.lower()).strip("-") or "x"

