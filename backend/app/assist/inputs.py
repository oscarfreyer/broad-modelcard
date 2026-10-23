"""The study the maintainer submits: a PDF, plain text, or a CSV/XLSX table.

Tables are rendered as text with row numbers (``row 3: …``) so the model can
cite a cell as "row 3, column Performance". Only the sha256 and size of the
input are ever reported or logged, never its content. A workbook's zip
directory is checked before openpyxl opens it (``STUDY_ARCHIVE_LIMITS``), so a
decompression bomb is refused instead of expanded.
"""

from __future__ import annotations

import csv
import hashlib
import io
from dataclasses import dataclass
from itertools import islice
from pathlib import Path
from typing import Iterable

from ..importers.spreadsheets import MB, ArchiveLimits, xlsx_archive_problem

PDF_SIGNATURE = b"%PDF-"
XLSX_SIGNATURE = b"PK\x03\x04"
MAX_INPUT_BYTES = 10 * 1024 * 1024
MAX_TEXT_CHARS = 200_000  # pasted text is refused above it; file-derived text is cut to it
MAX_TABLE_ROWS = 2_000  # per table or sheet
MAX_TABLE_COLUMNS = 100
MAX_SHEETS = 20
STUDY_ARCHIVE_LIMITS = ArchiveLimits(total=100 * MB, whole_part=100 * MB)
CUT_NOTE = f"(the study continues; only the first {MAX_TEXT_CHARS:,} characters are included)"
TEXT_SUFFIXES = (".txt", ".md", ".text")
TABLE_SUFFIXES = (".csv", ".tsv")


class InputError(ValueError):
    """A study the helper cannot read: (status, code, detail)."""

    def __init__(self, status: int, code: str, detail: str):
        super().__init__(detail)
        self.status = status
        self.code = code
        self.detail = detail


@dataclass(frozen=True)
class StudyInput:
    media: str  # pdf, text, csv, xlsx
    sha256: str
    size: int
    filename: str | None
    text: str | None = None  # everything but PDF
    pdf: bytes | None = None
    truncated: bool = False  # rows, columns, sheets, or characters were left out

    def info(self) -> dict:
        return {"media": self.media, "bytes": self.size, "sha256": self.sha256, "filename": self.filename, "truncated": self.truncated}


def digest(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()


def from_text(text: str) -> StudyInput:
    body = text.strip()
    if not body:
        raise InputError(422, "assist_unsupported_input", "The pasted text is empty.")
    if len(body) > MAX_TEXT_CHARS:
        raise InputError(422, "assist_unsupported_input", f"Pasted text is limited to {MAX_TEXT_CHARS:,} characters; upload the study as a file.")
    raw = body.encode("utf-8")
    return StudyInput("text", digest(raw), len(raw), None, text=body)


def from_file(filename: str | None, content: bytes) -> StudyInput:
    name = Path(filename or "study").name[:200]
    suffix = Path(name).suffix.lower()
    if not content:
        raise InputError(422, "assist_unsupported_input", "The file is empty.")
    if len(content) > MAX_INPUT_BYTES:
        raise InputError(413, "upload_too_large", f"The study file is larger than {MAX_INPUT_BYTES // (1024 * 1024)} MB, the most the helper accepts.")
    sha = digest(content)
    if suffix == ".pdf" or content.startswith(PDF_SIGNATURE):
        if not content.startswith(PDF_SIGNATURE):
            raise InputError(415, "unsupported_file_type", "The file is named .pdf but is not a PDF.")
        return StudyInput("pdf", sha, len(content), name, pdf=content)
    if suffix == ".xlsx":
        if not content.startswith(XLSX_SIGNATURE):
            raise InputError(415, "unsupported_file_type", "The file is named .xlsx but is not an Excel workbook.")
        text, cut = _xlsx_text(content, name)
        return _capped("xlsx", sha, len(content), name, text, cut)
    if suffix in TEXT_SUFFIXES + TABLE_SUFFIXES:
        if b"\x00" in content[:4096]:
            raise InputError(415, "unsupported_file_type", "The file is not text.")
        text = _decode(content)
        if suffix in TABLE_SUFFIXES:
            table, cut = _csv_text(text, name, "\t" if suffix == ".tsv" else None)
            return _capped("csv", sha, len(content), name, table, cut)
        return _capped("text", sha, len(content), name, text, False)
    raise InputError(415, "unsupported_file_type", "The helper reads PDF, text (.txt, .md), CSV, and XLSX files.")


def _capped(media: str, sha: str, size: int, name: str, text: str, cut: bool) -> StudyInput:
    """The file's text, cut to the pasted-text limit (and said so, to the model and in the response)."""
    if len(text) > MAX_TEXT_CHARS:
        text, cut = text[:MAX_TEXT_CHARS].rstrip() + "\n" + CUT_NOTE, True
    return StudyInput(media, sha, size, name, text=text, truncated=cut)


def _decode(content: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    return content.decode("latin-1")


def _rows_as_text(title: str, rows: Iterable[Iterable[object]]) -> tuple[str, bool]:
    """(text, cut): at most MAX_TABLE_ROWS rows of at most MAX_TABLE_COLUMNS cells,
    numbered as in the file; stops reading once the text passes MAX_TEXT_CHARS."""
    lines, length, cut = [title], len(title), False
    for number, cells in enumerate(rows, start=1):
        if number > MAX_TABLE_ROWS:
            lines.append(f"(the table continues; only the first {MAX_TABLE_ROWS:,} rows are included)")
            return "\n".join(lines), True
        values = ["" if cell is None else str(cell).strip() for cell in cells]
        if len(values) > MAX_TABLE_COLUMNS:
            cut = cut or any(values[MAX_TABLE_COLUMNS:])
            values = values[:MAX_TABLE_COLUMNS]
        if any(values):
            line = f"row {number}: " + " | ".join(values)
            lines.append(line)
            length += len(line) + 1
            if length > MAX_TEXT_CHARS:
                return "\n".join(lines), True  # cut to the limit by the caller
    if cut:
        lines.append(f"(the table has more columns; only the first {MAX_TABLE_COLUMNS} are included)")
    return "\n".join(lines), cut


def _csv_text(text: str, name: str, delimiter: str | None) -> tuple[str, bool]:
    if delimiter is None:
        try:
            delimiter = csv.Sniffer().sniff(text[:4096], delimiters=",;\t").delimiter
        except csv.Error:
            delimiter = ","
    try:
        # Read row by row: stops at the row and character limits, never parses the rest.
        return _rows_as_text(f"Table {name} (row numbers as in the file; cells separated by |):", csv.reader(io.StringIO(text), delimiter=delimiter))
    except csv.Error as error:  # e.g. a cell longer than the csv module's field limit
        raise InputError(422, "assist_unsupported_input", f"The table could not be read as CSV ({error}). Paste the part with the results instead.") from None


def _xlsx_text(content: bytes, name: str) -> tuple[str, bool]:
    from openpyxl import load_workbook  # already a backend dependency (uploads)

    problem = xlsx_archive_problem(io.BytesIO(content), STUDY_ARCHIVE_LIMITS, unreadable="The workbook could not be read.")
    if problem is not None:
        raise InputError(422, "assist_unsupported_input", problem)
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as error:  # noqa: BLE001 (any unreadable workbook is the same answer)
        raise InputError(422, "assist_unsupported_input", "The workbook could not be read.") from error
    try:
        parts, cut = [], len(workbook.sheetnames) > MAX_SHEETS
        for sheet in islice(workbook.worksheets, MAX_SHEETS):
            # Streamed: one row or column past the limits tells that the sheet continues.
            rows = sheet.iter_rows(values_only=True, max_row=MAX_TABLE_ROWS + 1, max_col=MAX_TABLE_COLUMNS + 1)
            text, sheet_cut = _rows_as_text(f"Sheet {sheet.title} of {name} (row numbers as in the sheet; cells separated by |):", rows)
            parts.append(text)
            cut = cut or sheet_cut
            if sum(len(part) + 2 for part in parts) > MAX_TEXT_CHARS:
                return "\n\n".join(parts), True  # cut to the limit by the caller
        if len(workbook.sheetnames) > MAX_SHEETS:
            parts.append(f"(the workbook has more sheets; only the first {MAX_SHEETS} are included)")
        return "\n\n".join(parts), cut
    except InputError:
        raise
    except Exception as error:  # noqa: BLE001 (zipfile.BadZipFile, KeyError, ValueError, TypeError, zlib errors: unreadable)
        raise InputError(422, "assist_unsupported_input", "The workbook could not be read.") from error
    finally:
        workbook.close()
