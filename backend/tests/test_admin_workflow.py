"""The update workflow through the API on a private tiny store: data sets,
templates, staged reads on every read endpoint, change summaries, publish,
rollback, the new-model flow, and request limits."""

from __future__ import annotations

import io

import openpyxl
import pytest
from admin_helpers import (
    ADMIN,
    TINY_TARGET,
    errors_by_code,
    notes_by_code,
    post_upload,
    read_csv,
    template,
    upload_and_wait,
    write_csv,
)
from conftest import local_client

from app.main import API_PREFIX, create_app

TINY = {"model": "Tiny-1", "use_case": "CDSS"}
ALL_1A00 = 0  # filled metrics rows: 1A00's 90 groups, then 1A01's
ALL_1A01 = 90


def edited(client, kind, edit, **params) -> bytes:
    header, rows, comments = read_csv(template(client, kind, filled="true", **params))
    result = edit(header, rows)
    if result:
        header, rows = result
    return write_csv(header, rows, comments)


def set_values(index: int, **values):
    names = {"perf": "Performance", "ci_low": "CI low", "ci_high": "CI high"}

    def edit(header, rows):
        for key, value in values.items():
            rows[index][header.index(names.get(key, key))] = value

    return edit


def stage_metrics(client, edit, **form) -> dict:
    upload = upload_and_wait(client, "metrics", edited(client, "metrics", edit, **TINY_TARGET), **TINY_TARGET, **form)
    assert upload["status"] == "valid", upload["report"]
    return upload


def get(client, path, **params):
    return client.get(API_PREFIX + path, params=params)


# ---------------------------------------------------------------- data sets

def test_datasets_list_the_ten_data_sets(admin_client):
    datasets = {info["kind"]: info for info in admin_client.get(f"{ADMIN}/datasets").json()}
    assert list(datasets) == [
        "models", "metrics", "multimorbidity", "prevalence", "safety", "population", "studies", "effects", "local_metrics", "model_cards"
    ]
    assert all(info["version"] == 1 and info["staged_upload_id"] is None for info in datasets.values())
    assert datasets["models"]["counts"] == {"models": 1, "use_cases": 1}
    assert datasets["metrics"]["counts"] == {"use_cases": 1, "metric_rows": 180}
    assert datasets["multimorbidity"]["counts"] == {"use_cases": 1, "combinations": 1, "metric_rows": 1}
    assert datasets["prevalence"]["counts"] == {"codes": 3}
    assert datasets["safety"]["counts"] == {"incidents": 0, "open": 0}
    assert datasets["population"]["counts"] == {"populations": 0, "codes": 0}  # contract v1.10
    assert datasets["studies"]["counts"] == {"studies": 0}  # contract v1.13
    assert datasets["local_metrics"]["counts"] == {"use_cases": 0, "populations": 0, "metric_rows": 0}  # §26
    assert datasets["model_cards"]["counts"] == {"model_cards": 0}  # §26
    assert datasets["effects"]["counts"] == {"effects": 0}  # contract v1.13, §27
    upload = stage_metrics(admin_client, set_values(ALL_1A01, perf="0.62"))
    datasets = {info["kind"]: info for info in admin_client.get(f"{ADMIN}/datasets").json()}
    assert datasets["metrics"]["staged_upload_id"] == upload["upload_id"]
    assert datasets["metrics"]["staged_upload_ids"] == [upload["upload_id"]]


def test_blank_templates(admin_client):
    header, rows, comments = read_csv(template(admin_client, "metrics", "csv"))
    assert rows == [] and header[:6] == ["ICD-11 code", "Condition", "Age", "Sex", "Ethnicity", "Setting"] and len(header) == 23
    assert header[19:] == ["Action", "Source", "Study", "Source date"]  # contract v1.7, optional
    assert any("Never upload patient-level or real patient data" in line for line in comments)
    workbook = openpyxl.load_workbook(io.BytesIO(template(admin_client, "metrics", "xlsx")))
    assert workbook.sheetnames == ["Instructions", "Data", "Lists"]
    assert workbook["Lists"].sheet_state == "hidden"
    instructions = "\n".join(str(cell.value) for row in workbook["Instructions"].iter_rows() for cell in row if cell.value)
    assert "Never upload patient-level or real patient data" in instructions
    assert "synthetic" in instructions.lower()
    data = workbook["Data"]
    assert [cell.value for cell in data[1]] == header and data.max_row == 1
    validated = {str(validation.sqref).split(":")[0] for validation in data.data_validations.dataValidation}
    letters = {name: openpyxl.utils.get_column_letter(index + 1) for index, name in enumerate(header)}
    for column in ("Age", "Sex", "Ethnicity", "Evidence grade", "Validation type", "In intended use", "Intended-use level"):
        assert f"{letters[column]}2" in validated, column
    lists = [[cell.value for cell in column if cell.value] for column in workbook["Lists"].iter_cols()]
    assert ["In scope", "Partially in scope", "Out of scope"] in lists
    assert ["All ages", "Age 0-17", "Age 18-39", "Age 40-64", "Age 65+"] in lists
    prevalence = openpyxl.load_workbook(io.BytesIO(template(admin_client, "prevalence", "xlsx")))
    tiers = [[cell.value for cell in column if cell.value] for column in prevalence["Lists"].iter_cols()]
    assert tiers == [["Add or replace", "Remove"], ["Very rare", "Rare", "Uncommon", "Common", "Very common"]]


@pytest.mark.parametrize("kind", ["models", "prevalence", "multimorbidity", "safety"])
def test_every_kind_has_both_templates(admin_client, kind):
    assert template(admin_client, kind, "csv").startswith(b"# Broad Modelcard")
    assert template(admin_client, kind, "xlsx").startswith(b"PK\x03\x04")


def test_template_downloads_are_attachments(admin_client):
    response = admin_client.get(f"{ADMIN}/templates/metrics.xlsx", params={"filled": "true", **TINY_TARGET})
    assert response.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert 'filename="broad-modelcard-metrics-tiny-1-cdss-v1.xlsx"' in response.headers["content-disposition"]
    assert response.headers["cache-control"] == "no-store"


# ------------------------------------------------------------ staged reads

def read_all(client, **params):
    """One request per read operation (13), keyed by name."""
    concept = dict(TINY, **params)
    return {
        "meta": get(client, "/meta", **params),
        "models": get(client, "/models", **params),
        "catalogue": get(client, "/catalogue", **params),
        "overview": get(client, "/atlas/overview", **params),
        "chapter": get(client, "/atlas/chapters/CH-01", **params),
        "map": get(client, "/metrics/map", scope="atlas", **concept),
        "summary": get(client, "/metrics/summary", view="performance", **concept),
        "concept": get(client, "/metrics/concepts/1A01", **concept),
        "candidates": get(client, "/multimorbidity/candidates", scope="atlas", **concept),
        "combination": get(client, "/multimorbidity/combination", selected="1A00,1A01", **concept),
        "search": get(client, "/search", q="1A0", **params),
        "export_concept": get(client, "/export/concepts/1A01", **concept),
        "export_dataset": get(client, "/export/dataset", **concept),
    }


def test_staged_reads_layer_the_upload_over_live_data_on_every_read_endpoint(admin_client):
    client = admin_client
    live_before = read_all(client)
    assert all(response.status_code == 200 for response in live_before.values())

    metrics = stage_metrics(client, set_values(ALL_1A01, perf="0.85", ci_low="0.8", ci_high="0.9"))
    multimorbidity = upload_and_wait(
        client,
        "multimorbidity",
        edited(client, "multimorbidity", set_values(0, perf="0.95", ci_low="0.9", ci_high="0.99"), **TINY_TARGET),
        **TINY_TARGET,
    )
    registry = upload_and_wait(client, "models", edited(client, "models", set_values(0, **{"Display name": "Tiny One"})))
    prevalence = upload_and_wait(client, "prevalence", edited(client, "prevalence", set_values(2, Commonness="Very common")))
    assert {u["status"] for u in (multimorbidity, registry, prevalence)} == {"valid"}
    assert not client.get(f"{ADMIN}/uploads/{metrics['upload_id']}").json()["previewed"]

    staged = {name: read_all(client, version=f"staged:{upload['upload_id']}") for name, upload in (
        ("metrics", metrics), ("multimorbidity", multimorbidity), ("models", registry), ("prevalence", prevalence),
    )}
    for name, responses in staged.items():
        upload_id = {"metrics": metrics, "multimorbidity": multimorbidity, "models": registry, "prevalence": prevalence}[name]["upload_id"]
        for endpoint, response in responses.items():
            assert response.status_code == 200, (name, endpoint, response.text[:200])
            assert response.headers["X-Dataset-Version"] == f"staged:{upload_id}", (name, endpoint)
        assert responses["meta"].json()["dataset_version"] == f"staged:{upload_id}"
        assert client.get(f"{ADMIN}/uploads/{upload_id}").json()["previewed"] is True

    # The staged content shows where it should.
    assert staged["metrics"]["concept"].json()["metrics"]["perf"] == 0.85
    assert live_before["concept"].json()["metrics"]["perf"] == 0.6
    map_rows = {row[0]: row for row in staged["metrics"]["map"].json()["rows"]}
    columns = staged["metrics"]["map"].json()["columns"]
    assert map_rows["1A01"][columns.index("perf")] == 0.85
    assert staged["metrics"]["summary"].json()["counts"] != live_before["summary"].json()["counts"]
    assert staged["metrics"]["export_concept"].json()["model_card"]["perf"] == 0.85
    assert '"perf":0.85' in staged["metrics"]["export_dataset"].text
    assert staged["multimorbidity"]["combination"].json()["metrics"]["perf"] == 0.95
    assert staged["models"]["models"].json()["models"][0]["display_name"] == "Tiny One"
    assert staged["models"]["catalogue"].json()["tiles"][0]["system"]["display_name"] == "Tiny One"
    assert {name: responses["catalogue"].json()["tiles"][0]["staged_change"] for name, responses in staged.items()} == {
        "metrics": True, "multimorbidity": True, "models": True, "prevalence": False,
    }
    tiers = {concept["concept_id"]: concept["prevalence_tier"] for concept in staged["prevalence"]["overview"].json()["conditions"]}
    assert tiers["1A01"] == "very_common"

    # Live reads are unaffected, byte for byte.
    live_after = read_all(client)
    for endpoint in live_before:
        if endpoint == "export_concept":  # carries its own timestamp
            continue
        assert live_after[endpoint].headers["ETag"] == live_before[endpoint].headers["ETag"], endpoint
        assert live_after[endpoint].headers["X-Dataset-Version"] == "1"


def test_staged_read_errors(admin_client):
    client = admin_client
    response = get(client, "/meta", version="staged:up-unknown")
    assert response.status_code == 404 and response.json()["code"] == "unknown_upload"
    invalid = upload_and_wait(client, "prevalence", b"ICD-11 code,Commonness\nZZ99,Rare\n")
    assert invalid["status"] == "invalid"
    response = get(client, "/meta", version=f"staged:{invalid['upload_id']}")
    assert response.status_code == 409 and response.json()["code"] == "upload_not_valid"
    valid = stage_metrics(client, set_values(ALL_1A01, perf="0.61"))
    assert get(client, "/models", version=f"staged:{valid['upload_id']}").status_code == 200
    client.post(f"{ADMIN}/uploads/{valid['upload_id']}/publish", json={"previewed": True})
    response = get(client, "/models", version=f"staged:{valid['upload_id']}")
    assert response.status_code == 409 and response.json()["code"] == "upload_not_valid"


def test_staged_and_live_responses_never_share_cache_entries(admin_client):
    client = admin_client
    live = get(client, "/metrics/map", scope="atlas", **TINY)
    upload = stage_metrics(client, set_values(ALL_1A01, perf="0.9", ci_low="0.85", ci_high="0.95"))
    staged = get(client, "/metrics/map", scope="atlas", version=f"staged:{upload['upload_id']}", **TINY)
    again = get(client, "/metrics/map", scope="atlas", **TINY)
    assert staged.headers["ETag"] != live.headers["ETag"]
    assert again.headers["ETag"] == live.headers["ETag"] and again.content == live.content
    revalidated = client.get(f"{API_PREFIX}/metrics/map", params={"scope": "atlas", **TINY, "version": f"staged:{upload['upload_id']}"},
                             headers={"If-None-Match": staged.headers["ETag"]})
    assert revalidated.status_code == 304 and revalidated.headers["X-Dataset-Version"] == f"staged:{upload['upload_id']}"


def test_staged_reads_follow_a_check_run_again_on_the_same_live_version(admin_client):
    """Review fix: a staged upload checked again writes a new content set under
    the same upload id (and deletes the old one). A staged read made just
    before (here: between a publish and the re-check it queues) used to keep
    serving the deleted set: every row read No data."""
    client = admin_client
    admin = client.app.state.api.admin
    upload = stage_metrics(client, set_values(ALL_1A01, perf="0.9", ci_low="0.85", ci_high="0.95"))
    staged = f"staged:{upload['upload_id']}"
    first = get(client, "/metrics/concepts/1A00", version=staged, **TINY).json()["metrics"]
    assert first is not None
    set_id = admin._require(upload["upload_id"]).set_id
    # Another staged set above it, so the new set cannot take the freed id again.
    other = client.post(f"{ADMIN}/uploads/records", json={"kind": "prevalence", "records": [
        {"code": "1A00", "tier": "Common", "source": "Test", "study": "T-1", "source_date": "2026-10-01"}]})
    assert other.status_code == 202, other.text
    admin.wait_for_checks()
    admin._recheck([(upload["upload_id"], set_id)])  # what a publish queues for dependent uploads
    admin.wait_for_checks()
    assert admin._require(upload["upload_id"]).set_id != set_id
    # A lens the first read did not load, so its rows come from the store now.
    again = get(client, "/metrics/concepts/1A01", version=staged, age="Age 65+", **TINY).json()["metrics"]
    assert again is not None and again["provenance"]["source"] == "condition"
    summary = client.get(f"{ADMIN}/uploads/{upload['upload_id']}/changes")
    assert summary.status_code == 200 and summary.json()["rows"]["changed"] == 1


# ---------------------------------------------------------- change summary

def counts(rows: dict) -> dict:
    """Row counts without the v1.4 ``samples``."""
    return {name: value for name, value in rows.items() if name != "samples"}


def names(samples: list[dict]) -> list[tuple]:
    return [(sample["key"], sample["name"], sample["code"]) for sample in samples]


def test_change_summary_reports_rows_status_shifts_and_removals(admin_client):
    client = admin_client

    def edit(header, rows):
        rows[ALL_1A01][header.index("Performance")] = "0.85"
        rows[ALL_1A01][header.index("CI low")] = "0.8"
        rows[ALL_1A01][header.index("CI high")] = "0.9"
        del rows[ALL_1A00:ALL_1A00 + 90]  # 1A00 loses all its data

    upload = stage_metrics(client, edit)
    summary = client.get(f"{ADMIN}/uploads/{upload['upload_id']}/changes").json()
    assert summary["base_version"] == 1 and summary["kind"] == "metrics"
    assert counts(summary["rows"]) == {"added": 0, "changed": 1, "removed": 90, "unchanged": 89}
    shifts = {(shift["view"], shift["from"], shift["to"]): shift for shift in summary["status_shifts"]}
    assert shifts[("performance", "weak", "strong")]["sample_concept_ids"] == ["1A01"]
    assert shifts[("performance", "strong", "no_data")]["sample_concept_ids"] == ["1A00"]
    assert shifts[("intended_use", "in_scope", "no_data")]["count"] == 1
    assert all(shift["count"] == len(shift["sample_concept_ids"]) for shift in summary["status_shifts"])

    # v1.4: every shift names its one target, and the rows are named.
    for shift in summary["status_shifts"]:
        assert shift["targets"] == [
            {**TINY_TARGET, "count": shift["count"], "sample_concept_ids": shift["sample_concept_ids"]}
        ]
    samples = summary["rows"]["samples"]
    assert samples["added"] == []
    assert [(sample["concept_id"], sample["lens"]) for sample in samples["changed"]] == [
        ("1A01", {"age": "All ages", "sex": "All sexes", "ethnicity": "All ethnicities"})
    ]
    assert samples["changed"][0]["name"] and samples["changed"][0]["code"] == "1A01"
    assert len(samples["removed"]) == 50  # capped; the count says 90
    assert {sample["concept_id"] for sample in samples["removed"]} == {"1A00"}
    assert len({sample["key"] for sample in samples["removed"]}) == 50


def test_registry_change_summary(admin_client):
    client = admin_client
    upload = upload_and_wait(client, "models", edited(client, "models", set_values(0, Status="EXPERIMENTAL")))
    summary = client.get(f"{ADMIN}/uploads/{upload['upload_id']}/changes").json()
    assert counts(summary["rows"]) == {"added": 0, "changed": 1, "removed": 0, "unchanged": 0}
    assert names(summary["rows"]["samples"]["changed"]) == [("Tiny-1 / CDSS", "Tiny-1 / CDSS", None)]
    assert summary["models"] == {"added": [], "changed": ["Tiny-1"], "removed": []}
    assert summary["use_cases"] == {"added": [], "removed": []}
    assert summary["status_shifts"] == []


def test_prevalence_change_summary_counts_codes(admin_client):
    client = admin_client

    def edit(header, rows):
        rows[2][header.index("Commonness")] = "Very common"  # 1A01
        del rows[1]  # 1A00.0

    upload = upload_and_wait(client, "prevalence", edited(client, "prevalence", edit))
    summary = client.get(f"{ADMIN}/uploads/{upload['upload_id']}/changes").json()
    assert counts(summary["rows"]) == {"added": 0, "changed": 1, "removed": 1, "unchanged": 1}
    samples = summary["rows"]["samples"]
    assert [sample["code"] for sample in samples["changed"]] == ["1A01"]
    assert [sample["code"] for sample in samples["removed"]] == ["1A00.0"]
    assert all(sample["lens"] is None and sample["name"] != sample["code"] for sample in samples["changed"] + samples["removed"])


# ---------------------------------------------------------- publish/rollback

def served_rows(client) -> dict:
    """What the atlas shows for Tiny-1, every view and some lenses, without version labels."""
    out = {}
    for lens in ({}, {"age": "Age 65+"}, {"sex": "Sex Female", "ethnicity": "Ethnicity Asian"}):
        body = get(client, "/metrics/map", scope="atlas", **TINY, **lens).json()
        out[("map", tuple(lens.items()))] = body["rows"]
        for concept in ("CH-01", "BLK-01", "1A00", "1A01", "1A00.0"):
            out[(concept, tuple(lens.items()))] = get(client, f"/metrics/concepts/{concept}", **TINY, **lens).json()["metrics"]
        for view in ("overall", "performance", "evidence", "intended_use"):
            summary = get(client, "/metrics/summary", view=view, **TINY, **lens).json()
            out[("summary", view, tuple(lens.items()))] = (summary["counts"], summary["chapters"])
    out["combination"] = get(client, "/multimorbidity/combination", selected="1A00,1A01", **TINY).json()["metrics"]
    out["models"] = get(client, "/models").json()
    out["overview"] = get(client, "/atlas/overview").json()["conditions"]
    return out


def test_publish_then_rollback_restores_every_served_row_exactly(admin_client):
    client = admin_client
    before = served_rows(client)
    upload = stage_metrics(client, set_values(ALL_1A01, perf="0.85", ci_low="0.8", ci_high="0.9"))
    published = client.post(f"{ADMIN}/uploads/{upload['upload_id']}/publish", json={"previewed": False}).json()
    assert published["version_id"] == 2 and published["previewed"] is False
    assert "Performance metrics for Tiny-1 / CDSS" in published["summary"]
    assert get(client, "/meta").headers["X-Dataset-Version"] == "2"
    assert client.get(f"{ADMIN}/uploads/{upload['upload_id']}").json()["status"] == "published"
    changed = served_rows(client)
    assert changed != before
    rolled = client.post(f"{ADMIN}/versions/1/rollback", params={"kind": "metrics"}).json()
    assert (rolled["version_id"], rolled["kind"], rolled["rollback_of"]) == (3, "metrics", 1)
    assert served_rows(client) == before
    history = client.get(f"{ADMIN}/versions").json()
    assert [entry["version_id"] for entry in history] == [3, 2, 1]
    assert [entry["version_id"] for entry in client.get(f"{ADMIN}/versions", params={"kind": "prevalence"}).json()] == [1]
    datasets = {info["kind"]: info for info in client.get(f"{ADMIN}/datasets").json()}
    assert datasets["metrics"]["version"] == 3 and datasets["models"]["version"] == 1


def test_rollback_of_the_import_without_kind_restores_every_data_set(admin_client):
    client = admin_client
    before = served_rows(client)
    for kind, content in (
        ("prevalence", edited(client, "prevalence", set_values(2, Commonness="Common"))),
        ("models", edited(client, "models", set_values(0, **{"Display name": "Renamed"}))),
    ):
        upload = upload_and_wait(client, kind, content)
        assert client.post(f"{ADMIN}/uploads/{upload['upload_id']}/publish", json={"previewed": True}).status_code == 200
    assert served_rows(client) != before
    rolled = client.post(f"{ADMIN}/versions/1/rollback").json()
    assert rolled["kind"] is None and rolled["summary"] == "Rolled back all data sets to version 1."
    assert served_rows(client) == before


def test_previewed_is_recorded_when_the_atlas_previewed_the_upload(admin_client):
    client = admin_client
    upload = stage_metrics(client, set_values(ALL_1A01, perf="0.62"))
    get(client, "/meta", version=f"staged:{upload['upload_id']}")
    entry = client.post(f"{ADMIN}/uploads/{upload['upload_id']}/publish", json={"previewed": False}).json()
    assert entry["previewed"] is True


def test_publish_and_discard_need_a_staged_valid_upload(admin_client):
    client = admin_client
    invalid = upload_and_wait(client, "prevalence", b"ICD-11 code,Commonness\nZZ99,Rare\n")
    response = client.post(f"{ADMIN}/uploads/{invalid['upload_id']}/publish", json={"previewed": False})
    assert response.status_code == 409 and response.json()["code"] == "upload_not_valid"
    assert client.post(f"{ADMIN}/uploads/up-nothing/publish", json={"previewed": False}).status_code == 404
    assert client.post(f"{ADMIN}/uploads/{invalid['upload_id']}/publish", json={}).status_code == 422
    exists = post_upload(client, "prevalence", b"ICD-11 code,Commonness\n1A00,Rare\n")
    assert exists.status_code == 409 and exists.json()["code"] == "staged_upload_exists"
    replaced = upload_and_wait(client, "prevalence", b"ICD-11 code,Commonness\n1A00,Rare\n1A01,Rare\n", replace_staged=True)
    assert client.get(f"{ADMIN}/uploads/{invalid['upload_id']}").json()["status"] == "superseded"
    assert replaced["status"] == "valid"
    assert client.delete(f"{ADMIN}/uploads/{replaced['upload_id']}").status_code == 204
    assert client.get(f"{ADMIN}/uploads/{replaced['upload_id']}").json()["status"] == "discarded"
    assert [entry["version_id"] for entry in client.get(f"{ADMIN}/versions").json()] == [1]


# ------------------------------------------------------------ new models

def test_a_new_model_is_added_metrics_first(admin_client):
    client = admin_client

    def add_model(header, rows):
        rows.append(["Tiny-2", "Tiny-2", "Tiny", "2", "EXPERIMENTAL", "CDSS", "Yes", "No"])

    registry = upload_and_wait(client, "models", edited(client, "models", add_model))
    assert list(errors_by_code(registry)) == ["use_case_without_metrics"]

    new_target = {"model_version": "Tiny-2", "use_case": "CDSS"}
    metrics = upload_and_wait(client, "metrics", template(client, "metrics", filled="true", **TINY_TARGET), **new_target)
    assert metrics["status"] == "valid" and "target_not_published" in notes_by_code(metrics)
    summary = client.get(f"{ADMIN}/uploads/{metrics['upload_id']}/changes").json()
    assert summary["rows"]["added"] == 180
    assert client.post(f"{ADMIN}/uploads/{metrics['upload_id']}/publish", json={"previewed": False}).status_code == 200
    assert get(client, "/metrics/map", scope="atlas", model="Tiny-2", use_case="CDSS").status_code == 404  # not listed yet

    client.app.state.api.admin.wait_for_checks()  # the staged registry upload is checked again
    registry = client.get(f"{ADMIN}/uploads/{registry['upload_id']}").json()
    assert registry["status"] == "valid", registry["report"]
    changes = client.get(f"{ADMIN}/uploads/{registry['upload_id']}/changes").json()
    assert changes["models"]["added"] == ["Tiny-2"] and changes["use_cases"]["added"] == ["Tiny-2 / CDSS"]
    assert client.post(f"{ADMIN}/uploads/{registry['upload_id']}/publish", json={"previewed": True}).status_code == 200
    assert [model["model_version"] for model in get(client, "/models").json()["models"]] == ["Tiny-1", "Tiny-2"]
    assert get(client, "/metrics/map", scope="atlas", model="Tiny-2", use_case="CDSS").status_code == 200

    # Rolling the metrics back to the import would leave Tiny-2 without metrics.
    response = client.post(f"{ADMIN}/versions/1/rollback", params={"kind": "metrics"})
    assert response.status_code == 409 and response.json()["code"] == "rollback_conflict"
    assert client.post(f"{ADMIN}/versions/1/rollback", params={"kind": "models"}).status_code == 200
    assert client.post(f"{ADMIN}/versions/1/rollback", params={"kind": "metrics"}).status_code == 200


def test_metrics_for_a_target_that_left_the_staged_registry_are_outdated(admin_client):
    client = admin_client

    def add_model(header, rows):
        rows.append(["Tiny-3", "Tiny-3", "", "", "ACTIVE", "CDSS", "No", "No"])

    registry = upload_and_wait(client, "models", edited(client, "models", add_model))
    metrics = upload_and_wait(client, "metrics", template(client, "metrics", filled="true", **TINY_TARGET), model_version="Tiny-3", use_case="CDSS")
    assert metrics["status"] == "valid"
    client.delete(f"{ADMIN}/uploads/{registry['upload_id']}")
    response = client.post(f"{ADMIN}/uploads/{metrics['upload_id']}/publish", json={"previewed": False})
    assert response.status_code == 409 and response.json()["code"] == "upload_outdated"


# --------------------------------------------------------------- limits

def test_oversized_uploads_get_413(tiny_store, tmp_path):
    app = create_app(tiny_store, max_upload_bytes=2_000)
    with local_client(app) as client:
        big = b"ICD-11 code,Commonness\n" + b"1A00,Rare\n" * 20_000
        response = post_upload(client, "prevalence", big)
        assert response.status_code == 413 and response.json()["code"] == "upload_too_large"
        assert response.headers["content-type"].startswith("application/problem+json")

        boundary = "bmcboundary"
        body = (
            f"--{boundary}\r\nContent-Disposition: form-data; name=\"kind\"\r\n\r\nprevalence\r\n"
            f"--{boundary}\r\nContent-Disposition: form-data; name=\"file\"; filename=\"p.csv\"\r\n"
            "Content-Type: text/csv\r\n\r\n"
        ).encode() + big + f"\r\n--{boundary}--\r\n".encode()

        def chunks():  # no Content-Length: the size is counted while reading
            for start in range(0, len(body), 4096):
                yield body[start:start + 4096]

        response = client.post(
            f"{ADMIN}/uploads", content=chunks(), headers={"Content-Type": f"multipart/form-data; boundary={boundary}"}
        )
        assert response.status_code == 413 and response.json()["code"] == "upload_too_large"
        small = post_upload(client, "prevalence", b"ICD-11 code,Commonness\n1A00,Rare\n")
        assert small.status_code == 202
    assert not list((tiny_store.parent / "uploads" / ".incoming").iterdir())


def test_every_other_request_body_is_bounded(admin_client):
    """Review fix: bodies of the small JSON requests (submissions, publish,
    rollback, the instance, ...) are bounded too (1 MB); before, any size was
    read into memory and parsed before validation refused it."""
    import json

    client = admin_client
    big = json.dumps({"developer_name": "D", "model_version": "Tiny-1", "note": "x" * (2 * 1024 * 1024)}).encode()
    for method, path in (("POST", "/submissions"), ("PUT", "/instance"), ("POST", "/versions/1/rollback")):
        response = client.request(method, f"{ADMIN}{path}", content=big, headers={"Content-Type": "application/json"})
        assert response.status_code == 413 and response.json()["code"] == "upload_too_large", (path, response.text[:200])

        def chunks():  # no Content-Length: the size is counted while reading
            for start in range(0, len(big), 65536):
                yield big[start:start + 65536]

        response = client.request(method, f"{ADMIN}{path}", content=chunks(), headers={"Content-Type": "application/json"})
        assert response.status_code == 413 and response.json()["code"] == "upload_too_large", (path, "chunked")
    ok = client.post(f"{ADMIN}/submissions", json={"developer_name": "D", "model_version": "Tiny-1"})
    assert ok.status_code == 201
