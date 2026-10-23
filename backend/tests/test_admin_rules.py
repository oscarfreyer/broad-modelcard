"""Validation rules through the API: each invalid upload is
rejected with row, column, and reason, nothing is published, and the
annotated file marks the cells. Runs on a private copy of the tiny store."""

from __future__ import annotations

import io

import openpyxl
import pytest
from admin_helpers import (
    ADMIN,
    TINY_TARGET,
    errors_by_code,
    notes_by_code,
    post_upload,
    read_csv,
    template,
    upload_and_wait,
    write_csv,
)

from app.publishing.files import ERROR_FILL, NOTE_FILL

FIRST_DATA_ROW = 4  # two comment lines, then the header on row 3


def metrics_csv(client, edit=None) -> bytes:
    header, rows, comments = read_csv(template(client, "metrics", filled="true", **TINY_TARGET))
    if edit:
        header, rows = edit(header, rows) or (header, rows)
    return write_csv(header, rows, comments)


def set_cell(column: str, value, index: int = 0):
    def edit(header, rows):
        rows[index][header.index(column)] = value

    return edit


def check_rejected(client, upload: dict, code: str, row: int | None, column: str | None, problem_part: str = ""):
    """The upload is invalid with this error; nothing was published."""
    assert upload["status"] == "invalid", upload
    group = errors_by_code(upload).get(code)
    assert group is not None, f"no {code} in {list(errors_by_code(upload))}"
    assert group["message"] and group["fix"] and group["count"] >= 1
    assert len(group["examples"]) <= 20
    example = group["examples"][0]
    if row is not None:
        assert example["row"] == row, example
    assert example["column"] == column, example
    assert problem_part in example["problem"], example
    versions = client.get(f"{ADMIN}/versions").json()
    assert [entry["version_id"] for entry in versions] == [1], "nothing may be published"
    return example


def check_marked(client, upload_id: str, row: int, column: str, problem_part: str = ""):
    """The annotated CSV lists the problem in that row; the XLSX highlights the cell and comments on it."""
    header, rows, _ = read_csv(client.get(f"{ADMIN}/uploads/{upload_id}/annotated.csv").content)
    assert header[0] == "Problems"
    marked = [cells for cells in rows if cells[0].startswith(f"Row {row}:")]
    assert marked and column in marked[0][0] and problem_part in marked[0][0], rows[:3]

    workbook = openpyxl.load_workbook(io.BytesIO(client.get(f"{ADMIN}/uploads/{upload_id}/annotated.xlsx").content))
    sheet = workbook["Data"]
    headers = [cell.value for cell in sheet[1]]
    target_row = next(
        cells for cells in sheet.iter_rows(min_row=2) if isinstance(cells[0].value, str) and cells[0].value.startswith(f"Row {row}:")
    )
    cell = target_row[headers.index(column)]
    assert cell.fill.fgColor.rgb.endswith(ERROR_FILL)
    assert cell.comment is not None and problem_part in cell.comment.text
    assert "Summary" in workbook.sheetnames


METRIC_CELL_CASES = [
    ("unknown_icd_code", "ICD-11 code", "ZZ99", "not an ICD-11 code"),
    ("unknown_icd_code", "ICD-11 code", "1A00.0", "subcategory"),
    ("unknown_icd_code", "ICD-11 code", "BLK-01", "chapter or block"),
    ("out_of_range", "Performance", "84", "between 0 and 1"),
    ("out_of_range", "Intended-use coverage", "1.5", "between 0 and 1"),
    ("out_of_range", "Samples", "-3", "at least 0"),
    ("not_a_number", "Performance", "high", "not a number"),
    ("not_a_number", "CI low", "84%", "not a number"),
    ("not_a_whole_number", "Studies", "2.5", "whole number"),
    ("unknown_evidence_grade", "Evidence grade", "G", "A, B, C, D, E, or F"),
    ("unknown_validation_type", "Validation type", "Randomised trial", "Unknown validation type"),
    ("unknown_intended_use_level", "Intended-use level", "Mostly", "In scope"),
    ("invalid_yes_no", "In intended use", "Maybe", "Yes or No"),
    ("invalid_warning_code", "Warnings", "Needs Review!", "not a warning code"),
    ("unknown_lens_value", "Age", "Age 100+", "Unknown age group"),
    ("unknown_lens_value", "Ethnicity", "Martian", "Unknown ethnicity group"),
    ("missing_value", "Intended-use level", "", "empty"),
    ("missing_value", "ICD-11 code", "", "empty"),
]


@pytest.mark.parametrize(("code", "column", "value", "problem"), METRIC_CELL_CASES, ids=[f"{c[0]}-{c[1]}-{c[2]}" for c in METRIC_CELL_CASES])
def test_metric_cell_errors(admin_client, code, column, value, problem):
    upload = upload_and_wait(admin_client, "metrics", metrics_csv(admin_client, set_cell(column, value, index=5)), **TINY_TARGET)
    check_rejected(admin_client, upload, code, FIRST_DATA_ROW + 5, column, problem)
    check_marked(admin_client, upload["upload_id"], FIRST_DATA_ROW + 5, column, problem)


def test_ci_low_above_ci_high(admin_client):
    def edit(header, rows):
        rows[2][header.index("CI low")] = "0.9"
        rows[2][header.index("CI high")] = "0.7"
        rows[2][header.index("Performance")] = "0.8"

    upload = upload_and_wait(admin_client, "metrics", metrics_csv(admin_client, edit), **TINY_TARGET)
    check_rejected(admin_client, upload, "ci_low_above_ci_high", FIRST_DATA_ROW + 2, "CI low", "above CI high")
    check_marked(admin_client, upload["upload_id"], FIRST_DATA_ROW + 2, "CI low", "above CI high")


def test_performance_outside_its_interval(admin_client):
    def edit(header, rows):
        rows[3][header.index("Performance")] = "0.99"
        rows[3][header.index("CI low")] = "0.7"
        rows[3][header.index("CI high")] = "0.8"

    upload = upload_and_wait(admin_client, "metrics", metrics_csv(admin_client, edit), **TINY_TARGET)
    check_rejected(admin_client, upload, "perf_outside_ci", FIRST_DATA_ROW + 3, "Performance", "outside its confidence interval")
    check_marked(admin_client, upload["upload_id"], FIRST_DATA_ROW + 3, "Performance")


def test_peer_reviewed_studies_cannot_exceed_studies(admin_client):
    def edit(header, rows):
        rows[0][header.index("Studies")] = "2"
        rows[0][header.index("Peer-reviewed studies")] = "3"

    upload = upload_and_wait(admin_client, "metrics", metrics_csv(admin_client, edit), **TINY_TARGET)
    check_rejected(admin_client, upload, "out_of_range", FIRST_DATA_ROW, "Peer-reviewed studies", "cannot exceed")


def test_duplicate_key(admin_client):
    def edit(header, rows):
        rows.insert(10, list(rows[3]))

    upload = upload_and_wait(admin_client, "metrics", metrics_csv(admin_client, edit), **TINY_TARGET)
    example = check_rejected(admin_client, upload, "duplicate_key", FIRST_DATA_ROW + 10, "ICD-11 code", "also appears in row")
    assert str(FIRST_DATA_ROW + 3) in example["problem"]
    check_marked(admin_client, upload["upload_id"], FIRST_DATA_ROW + 10, "ICD-11 code", "also appears")


def test_incomplete_lens_coverage(admin_client):
    def edit(header, rows):
        del rows[100]  # one of 1A01's 90 demographic groups

    upload = upload_and_wait(admin_client, "metrics", metrics_csv(admin_client, edit), **TINY_TARGET)
    example = check_rejected(admin_client, upload, "incomplete_lens_coverage", FIRST_DATA_ROW + 90, None, "has 89 of the 90")
    assert "missing" in example["problem"]


def test_missing_and_extra_columns(admin_client):
    def edit(header, rows):
        index = header.index("CI high")
        header = [name for position, name in enumerate(header) if position != index] + ["Notes"]
        rows = [[value for position, value in enumerate(row) if position != index] + ["x"] for row in rows]
        return header, rows

    upload = upload_and_wait(admin_client, "metrics", metrics_csv(admin_client, edit), **TINY_TARGET)
    check_rejected(admin_client, upload, "missing_column", 3, "CI high", "CI high is missing")
    extra = errors_by_code(upload)["extra_column"]["examples"][0]
    assert (extra["row"], extra["column"]) == (3, "Notes")
    _, _, comments = read_csv(admin_client.get(f"{ADMIN}/uploads/{upload['upload_id']}/annotated.csv").content)
    assert any("Header row: The column CI high is missing." in line for line in comments)


def test_wrong_template_for_the_data_set(admin_client):
    prevalence = template(admin_client, "prevalence", filled="true")
    upload = upload_and_wait(admin_client, "metrics", prevalence, **TINY_TARGET)
    example = check_rejected(admin_client, upload, "wrong_template", 0, None, "Prevalence template")
    assert "Prevalence" in errors_by_code(upload)["wrong_template"]["fix"]
    assert list(errors_by_code(upload)) == ["wrong_template"]


@pytest.mark.parametrize("content", [b"", b"# just a comment\n", None], ids=["zero-bytes", "comment-only", "header-only"])
def test_empty_file(admin_client, content):
    if content is None:
        header, _, _ = read_csv(template(admin_client, "metrics"))
        content = write_csv(header, [])
    upload = upload_and_wait(admin_client, "metrics", content, **TINY_TARGET)
    check_rejected(admin_client, upload, "empty_file", 0, None)


def test_unknown_model_or_use_case_for_the_target(admin_client):
    upload = upload_and_wait(admin_client, "metrics", metrics_csv(admin_client), model_version="Nope-1", use_case="CDSS")
    check_rejected(admin_client, upload, "unknown_target", 0, "model_version", "no AI system")
    upload = upload_and_wait(admin_client, "metrics", metrics_csv(admin_client), model_version="Tiny-1", use_case="Triage", replace_staged=True)
    check_rejected(admin_client, upload, "unknown_target", 0, "use_case", "no function")


def test_unreadable_workbook(admin_client):
    upload = upload_and_wait(admin_client, "metrics", b"PK\x03\x04 this is not a workbook", "broken.xlsx", **TINY_TARGET)
    check_rejected(admin_client, upload, "unreadable_file", 0, None, "not a readable Excel workbook")


# ------------------------------------------------------------- other kinds

def edited_template(client, kind, edit, **params):
    header, rows, comments = read_csv(template(client, kind, filled="true", **params))
    header, rows = edit(header, rows) or (header, rows)
    return write_csv(header, rows, comments)


def test_registry_use_case_without_metrics(admin_client):
    def edit(header, rows):
        rows.append([*rows[0][:5], "Triage", "No", "Yes"])

    upload = upload_and_wait(admin_client, "models", edited_template(admin_client, "models", edit))
    check_rejected(admin_client, upload, "use_case_without_metrics", FIRST_DATA_ROW + 1, "Function", "Tiny-1 / Triage")
    check_marked(admin_client, upload["upload_id"], FIRST_DATA_ROW + 1, "Function")


@pytest.mark.parametrize(
    ("column", "value", "code", "problem"),
    [
        ("Status", "BETA", "unknown_status", "ACTIVE or EXPERIMENTAL"),
        ("Default model", "Perhaps", "invalid_yes_no", "Yes or No"),
        ("Display name", "", "missing_value", "empty"),
        ("Model version", "x" * 80, "invalid_value", "longer than 64"),
    ],
)
def test_registry_cell_errors(admin_client, column, value, code, problem):
    def edit(header, rows):
        rows[0][header.index(column)] = value

    upload = upload_and_wait(admin_client, "models", edited_template(admin_client, "models", edit))
    check_rejected(admin_client, upload, code, FIRST_DATA_ROW, column, problem)


def test_registry_rows_of_one_model_must_agree_and_defaults_are_unique(admin_client):
    def edit(header, rows):
        rows.append([*rows[0]])
        rows[1][header.index("Function")] = "CDSS"
        rows[1][header.index("Display name")] = "Another name"
        rows.append(["Tiny-2", "Tiny-2", "", "", "ACTIVE", "CDSS", "Yes", "Yes"])

    upload = upload_and_wait(admin_client, "models", edited_template(admin_client, "models", edit))
    errors = errors_by_code(upload)
    assert "duplicate_key" in errors  # Tiny-1 / CDSS twice
    assert errors["invalid_default"]["examples"][0]["column"] == "Default model"


def test_prevalence_errors(admin_client):
    def edit(header, rows):
        rows[0][header.index("Commonness")] = "Frequent"
        rows.append(list(rows[1]))
        rows.append(["CH-01", "", "Rare"])

    upload = upload_and_wait(admin_client, "prevalence", edited_template(admin_client, "prevalence", edit))
    check_rejected(admin_client, upload, "unknown_tier", FIRST_DATA_ROW, "Commonness", "Very rare")
    assert errors_by_code(upload)["duplicate_key"]["examples"][0]["row"] == FIRST_DATA_ROW + 3
    assert "chapter or block" in errors_by_code(upload)["unknown_icd_code"]["examples"][0]["problem"]


@pytest.mark.parametrize(
    ("codes", "code", "problem"),
    [
        ("1A00", "invalid_condition_set", "2 to 12 conditions"),
        ("1A00 + 1A00", "invalid_condition_set", "appears twice"),
        ("1A00 + ZZ99", "unknown_icd_code", "not an ICD-11 code"),
    ],
)
def test_multimorbidity_condition_sets(admin_client, codes, code, problem):
    def edit(header, rows):
        rows[0][header.index("Condition codes")] = codes

    content = edited_template(admin_client, "multimorbidity", edit, **TINY_TARGET)
    upload = upload_and_wait(admin_client, "multimorbidity", content, **TINY_TARGET)
    check_rejected(admin_client, upload, code, FIRST_DATA_ROW, "Condition codes", problem)
    check_marked(admin_client, upload["upload_id"], FIRST_DATA_ROW, "Condition codes", problem)


# ------------------------------------------------------------------ notes

def test_notes_do_not_block(admin_client):
    def edit(header, rows):
        rows[0][header.index("Performance")] = "0.6"
        rows[0][header.index("CI low")] = "0.3"
        rows[0][header.index("CI high")] = "0.9"
        rows[1][header.index("Warnings")] = "brand_new_warning"
        rows[2][header.index("In intended use")] = "No"

    upload = upload_and_wait(admin_client, "metrics", metrics_csv(admin_client, edit), **TINY_TARGET)
    assert upload["status"] == "valid" and upload["report"]["errors"] == []
    notes = notes_by_code(upload)
    assert notes["very_wide_confidence_interval"]["examples"][0]["row"] == FIRST_DATA_ROW
    assert notes["unfamiliar_warning_code"]["examples"][0]["value"] == "brand_new_warning"
    assert notes["scope_mismatch"]["examples"][0]["row"] == FIRST_DATA_ROW + 2
    workbook = openpyxl.load_workbook(io.BytesIO(admin_client.get(f"{ADMIN}/uploads/{upload['upload_id']}/annotated.xlsx").content))
    sheet = workbook["Data"]
    headers = [cell.value for cell in sheet[1]]
    assert sheet.cell(row=2, column=headers.index("CI high") + 1).fill.fgColor.rgb.endswith(NOTE_FILL)


def test_many_errors_keep_twenty_examples_and_count_all(admin_client):
    def edit(header, rows):
        for row in rows:
            row[header.index("Evidence grade")] = "Z"

    upload = upload_and_wait(admin_client, "metrics", metrics_csv(admin_client, edit), **TINY_TARGET)
    group = errors_by_code(upload)["unknown_evidence_grade"]
    assert group["count"] == 180 and len(group["examples"]) == 20


def test_fixed_annotated_file_can_be_uploaded_again(admin_client):
    upload = upload_and_wait(admin_client, "metrics", metrics_csv(admin_client, set_cell("Performance", "high")), **TINY_TARGET)
    annotated = admin_client.get(f"{ADMIN}/uploads/{upload['upload_id']}/annotated.csv").content
    header, rows, comments = read_csv(annotated)
    rows[0][header.index("Performance")] = "0.8"
    fixed = upload_and_wait(admin_client, "metrics", write_csv(header, rows, comments), replace_staged=True, **TINY_TARGET)
    assert fixed["status"] == "valid", fixed["report"]


def test_semicolon_csv_with_decimal_commas(admin_client):
    header, rows, _ = read_csv(template(admin_client, "metrics", filled="true", **TINY_TARGET))
    text = "\n".join(";".join(cell.replace(".", ",") if cell.replace(".", "", 1).isdigit() else cell for cell in row) for row in [header, *rows])
    upload = upload_and_wait(admin_client, "metrics", text.encode("cp1252"), **TINY_TARGET)
    assert upload["status"] == "valid", upload["report"]
    changes = admin_client.get(f"{ADMIN}/uploads/{upload['upload_id']}/changes").json()
    assert changes["rows"]["changed"] == 0 and changes["rows"]["unchanged"] == 180


def test_upload_requests_are_checked(admin_client):
    assert post_upload(admin_client, "metrics", b"x", "a.csv").status_code == 422  # no target
    response = post_upload(admin_client, "prevalence", b"\xd0\xcf\x11\xe0legacy", "old.xls")
    assert response.status_code == 415 and response.json()["code"] == "unsupported_file_type"
    response = admin_client.post(f"{ADMIN}/uploads", data={"kind": "nonsense"}, files={"file": ("a.csv", b"x")})
    assert response.status_code == 422


def test_data_without_a_header_is_an_extra_column(admin_client):
    def edit(header, rows):
        rows[4].append("stray")

    upload = upload_and_wait(admin_client, "metrics", metrics_csv(admin_client, edit), **TINY_TARGET)
    example = check_rejected(admin_client, upload, "extra_column", FIRST_DATA_ROW + 4, None, "has data but no header")
    assert example["value"] == "stray" and "Column X" in example["problem"]  # after the 23 template columns (Setting since v1.14)


def test_upload_file_names_are_cleaned(admin_client):
    from app.publishing.uploads import clean_filename

    assert clean_filename('evil"\r\nname.csv') == "evil_name.csv"
    assert clean_filename("../../etc/passwd.csv") == "passwd.csv"
    assert clean_filename("") == "upload"
    upload = upload_and_wait(admin_client, "metrics", metrics_csv(admin_client), "../folder/metrics.csv", **TINY_TARGET)
    assert upload["filename"] == "metrics.csv"
    response = admin_client.get(f"{ADMIN}/uploads/{upload['upload_id']}/annotated.csv")
    assert response.status_code == 200 and 'filename="metrics-checked.csv"' in response.headers["content-disposition"]
