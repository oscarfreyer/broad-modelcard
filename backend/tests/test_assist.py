"""Assisted data entry (contract v1.9, CONTRACT.md §20).

Everything runs offline: the fake provider answers the fixed synthetic study
set, and a stub HTTP server on 127.0.0.1 plays the hosted (Anthropic Messages)
and local (OpenAI-compatible) providers. No real model is ever called.

The API key is a canary string; the tests assert it never appears in a
response body, a log record, the SQLite store, or the uploads folder.
"""

from __future__ import annotations

import base64
import copy
import io
import json
import logging
import os
import secrets
import stat
from pathlib import Path

import pytest
from admin_helpers import ADMIN, wait
from assist_stub import StubProvider, anthropic_text, anthropic_tool_use, openai_text
from conftest import copy_store, local_client

from app import db
from app.assist import config as configuration
from app.assist import evaluate
from app.assist.inputs import from_file, from_text
from app.assist.providers import FakeProvider
from app.assist.studyset import fake_answers, fake_provider, load_cases
from app.main import API_PREFIX, create_app

CANARY = "sk-ant-canary-" + secrets.token_hex(12)
CASES = {case.id: case for case in load_cases()}
ANSWERS = json.loads((db.BACKEND_DIR / "assist_eval" / "fake_answers.json").read_text())["answers"]


class _Collect(logging.Handler):
    def __init__(self):
        super().__init__(logging.DEBUG)
        self.lines: list[str] = []

    def emit(self, record: logging.LogRecord) -> None:
        self.lines.append(self.format(record))
        if record.exc_info:
            self.lines.append(logging.Formatter().formatException(record.exc_info))


@pytest.fixture(scope="module")
def logs():
    handler = _Collect()
    root = logging.getLogger()
    previous = root.level
    root.addHandler(handler)
    root.setLevel(logging.DEBUG)
    yield handler
    root.removeHandler(handler)
    root.setLevel(previous)


@pytest.fixture(scope="module")
def demo_store(demo_db, tmp_path_factory) -> Path:
    return copy_store(demo_db[0], tmp_path_factory.mktemp("assist") / "assist-demo.sqlite")


@pytest.fixture(scope="module")
def client(demo_store, logs):
    with local_client(create_app(demo_store, validate_responses=True)) as test_client:
        yield test_client


@pytest.fixture(scope="module")
def stub():
    with StubProvider() as server:
        yield server


def offline_provider_for(config):
    """The real providers, for the stub on 127.0.0.1 only: a test never reaches a hosted model."""
    from app.assist.providers import provider_for

    if not configuration.is_loopback_url(config.endpoint):
        pytest.fail(f"a helper test tried to reach {config.endpoint}; use the fake provider or the stub")
    return provider_for(config)


@pytest.fixture(autouse=True)
def helper_off(request):
    """Every test starts and ends with the helper off and the real providers, kept off the network."""
    if "client" not in request.fixturenames:
        yield
        return
    app_client = request.getfixturevalue("client")
    service = app_client.app.state.assist
    service.clear()
    service.provider_factory = offline_provider_for
    yield
    service.clear()
    service.provider_factory = offline_provider_for


def configure(client, **body):
    response = client.put(f"{ADMIN}/assist", json=body)
    assert response.status_code == 200, response.text
    return response.json()


def propose_file(client, case_id: str, **form):
    case = CASES[case_id]
    return client.post(
        f"{ADMIN}/assist/proposals",
        data={"kind": case.kind, **form},
        files={"file": (case.path.name, case.content(), "application/octet-stream")},
    )


def use_fake(client, kind: str = "hosted"):
    client.app.state.assist.provider_factory = lambda config: fake_provider(kind=config.provider)


# ------------------------------------------------------------- helper off

def test_off_by_default_explains_how_to_enable_and_the_rest_works(client):
    status = client.get(f"{ADMIN}/assist")
    assert status.status_code == 200 and status.headers["cache-control"] == "no-store"
    body = status.json()
    assert body["enabled"] is False and body["provider"] is None and body["key_set"] is False and body["accepts"] == []
    assert "API key" in body["how_to_enable"] and "local" in body["how_to_enable"]
    assert "only" in body["privacy"].lower() or "nowhere else" in body["privacy"]
    refused = client.post(f"{ADMIN}/assist/proposals", data={"kind": "metrics", "text": "Synthetic text."})
    assert refused.status_code == 409 and refused.json()["code"] == "assist_not_configured"
    assert "API key" in refused.json()["detail"]
    # The Data Library's own operations are unaffected.
    assert client.get(f"{ADMIN}/datasets").status_code == 200
    assert client.get(f"{API_PREFIX}/catalogue").status_code == 200


# ---------------------------------------------------------- configuration

def test_config_path_follows_the_store(tmp_path, monkeypatch):
    monkeypatch.delenv(configuration.CONFIG_ENV, raising=False)
    assert configuration.config_path(db.DEFAULT_DB_PATH) == db.DEFAULT_DB_PATH.parent / "assist.json"
    assert configuration.config_path(db.BACKEND_DIR / "data" / "dashboard-demo.sqlite").name == "dashboard-demo.assist.json"
    assert configuration.config_path(tmp_path / "store.sqlite") == tmp_path / "store.assist.json"
    monkeypatch.setenv(configuration.CONFIG_ENV, str(tmp_path / "elsewhere.json"))
    assert configuration.config_path(tmp_path / "store.sqlite") == tmp_path / "elsewhere.json"


def test_hosted_key_is_write_only_and_stored_with_mode_600(client):
    service = client.app.state.assist
    status = configure(client, provider="hosted", api_key=CANARY)
    assert status["enabled"] and status["provider"] == "hosted" and status["model"] == configuration.DEFAULT_HOSTED_MODEL
    assert status["key_set"] is True and status["base_url"] is None and "pdf" in status["accepts"]
    path = service.config_path
    assert path.name == "assist-demo.assist.json" and path.parent == Path(client.app.state.api.admin.db_path).parent
    assert stat.S_IMODE(os.stat(path).st_mode) == 0o600
    assert CANARY in path.read_text()  # the one place it lives
    for response in (client.get(f"{ADMIN}/assist"), client.put(f"{ADMIN}/assist", json={"provider": "hosted", "model": "claude-opus-5-5"})):
        assert response.status_code == 200 and CANARY not in response.text
    # Omitting the key keeps it; switching to a local model drops it.
    assert configuration.load(path).api_key == CANARY and configuration.load(path).model == "claude-opus-5-5"
    local = configure(client, provider="local", model="llama3.1:8b", base_url="http://127.0.0.1:11434/v1/")
    assert local["key_set"] is False and local["base_url"] == "http://127.0.0.1:11434/v1" and "pdf" not in local["accepts"]
    assert CANARY not in path.read_text()
    assert client.delete(f"{ADMIN}/assist").status_code == 204
    assert not path.exists() and client.get(f"{ADMIN}/assist").json()["enabled"] is False


@pytest.mark.parametrize(
    ("body", "fragment"),
    [
        ({"provider": "hosted"}, "API key"),
        ({"provider": "hosted", "api_key": CANARY + "\n"}, "printable"),
        ({"provider": "hosted", "api_key": "short"}, "printable"),
        ({"provider": "hosted", "api_key": CANARY, "base_url": "http://models.example/v1"}, "https"),
        ({"provider": "local", "model": "m"}, "base URL"),
        ({"provider": "local", "base_url": "http://127.0.0.1:1/v1"}, "model name"),
        ({"provider": "local", "model": "m", "base_url": f"http://user:{CANARY}@127.0.0.1:1/v1"}, "credentials"),
        ({"provider": "local", "model": "m", "base_url": "ftp://127.0.0.1/v1"}, "http"),
        ({"provider": "local", "model": "m", "base_url": "http://192.168.1.20:11434/v1"}, "this computer"),
        ({"provider": "local", "model": "m", "base_url": "https://models.example/v1", "api_key": CANARY}, "127.0.0.1, ::1, or localhost"),
        ({"provider": "local", "model": "m", "base_url": "http://localhost.example.org/v1"}, "this computer"),
        ({"provider": "remote", "api_key": CANARY}, "provider"),
        ({"provider": "hosted", "api_key": 12345678}, "api_key"),
    ],
)
def test_invalid_configuration_is_refused_without_echoing_values(client, body, fragment):
    response = client.put(f"{ADMIN}/assist", json=body)
    assert response.status_code == 422, response.text
    assert response.json()["code"] == "invalid_parameter" and fragment in response.json()["detail"]
    assert CANARY not in response.text and "12345678" not in response.text
    assert client.get(f"{ADMIN}/assist").json()["enabled"] is False


def test_a_stored_hosted_key_never_follows_a_changed_address(client):
    """Review S5 #1: changing only the base URL must not carry the stored key to the new host."""
    with StubProvider() as first, StubProvider() as second:
        first.respond = second.respond = lambda path, headers, body: (200, anthropic_text('{"records": [], "notes": []}'), {})
        configure(client, provider="hosted", api_key=CANARY, base_url=first.url)
        moved = client.put(f"{ADMIN}/assist", json={"provider": "hosted", "base_url": second.url})
        assert moved.status_code == 422 and moved.json()["code"] == "invalid_parameter"
        assert "address changed" in moved.json()["detail"] and "API key again" in moved.json()["detail"] and CANARY not in moved.text
        assert client.get(f"{ADMIN}/assist").json()["base_url"] == first.url  # the old configuration stays
        # The Anthropic API is another origin than the gateway: the key is asked for again.
        assert client.put(f"{ADMIN}/assist", json={"provider": "hosted"}).status_code == 422
        # The same origin with another path (or a trailing slash) keeps it.
        same = configure(client, provider="hosted", base_url=first.url + "/gateway/")
        assert same["key_set"] is True and same["base_url"] == first.url + "/gateway"
        # With the key typed again, the new address gets it, and only it.
        configure(client, provider="hosted", base_url=second.url, api_key=CANARY)
        response = client.post(f"{ADMIN}/assist/proposals", data={"kind": "metrics", "text": "Synthetic."})
        assert response.status_code == 200, response.text
        assert [request["headers"]["x-api-key"] for request in second.requests] == [CANARY]
        assert first.requests == []


def test_a_stored_local_key_is_dropped_when_the_address_changes(client):
    with StubProvider() as first, StubProvider() as second:
        first.respond = second.respond = lambda path, headers, body: (200, openai_text('{"records": [], "notes": []}'), {})
        configure(client, provider="local", model="m", api_key=CANARY, base_url=first.url + "/v1")
        assert configure(client, provider="local", model="m2", base_url=first.url + "/v1")["key_set"] is True  # same address
        moved = configure(client, provider="local", model="m", base_url=second.url + "/v1")
        assert moved["key_set"] is False
        assert CANARY not in client.app.state.assist.config_path.read_text()
        response = client.post(f"{ADMIN}/assist/proposals", data={"kind": "metrics", "text": "Synthetic."})
        assert response.status_code == 200, response.text
        assert len(second.requests) == 1 and "authorization" not in second.requests[0]["headers"]
        assert first.requests == []


def test_endpoint_origin_normalizes_scheme_host_and_port():
    origin = configuration.endpoint_origin
    assert origin("hosted", None) == origin("hosted", "https://API.anthropic.com:443/") == "https://api.anthropic.com:443"
    assert origin("local", "http://127.0.0.1:11434/v1") == origin("local", "http://127.0.0.1:11434/other")
    assert origin("local", "http://127.0.0.1:11434/v1") != origin("local", "http://127.0.0.1:11435/v1")
    assert origin("local", "http://localhost/v1") == "http://localhost:80"
    assert origin("local", "http://127.0.0.1/v1") != origin("local", "https://127.0.0.1/v1")


def test_a_local_model_is_on_this_computer(client, tmp_path):
    """Review S5 #6: a "local" base URL must be loopback, so studies and keys stay on this machine."""
    for url in ("http://127.0.0.1:11434/v1", "http://localhost:8080/v1", "http://[::1]:11434/v1"):
        assert configure(client, provider="local", model="m", base_url=url)["base_url"] == url
    # A configuration file from before the rule (or edited by hand) is not used.
    path = tmp_path / "old.assist.json"
    path.write_text(json.dumps({"provider": "local", "model": "m", "base_url": "http://10.0.0.5:11434/v1", "api_key": None}))
    assert configuration.load(path) is None


# ---------------------------------------------------------- the pipeline

def test_proposals_carry_sources_flags_and_only_existing_codes_and_cards(client):
    configure(client, provider="hosted", api_key=CANARY)
    use_fake(client)
    e1 = propose_file(client, "E1")
    assert e1.status_code == 200, e1.text
    body = e1.json()
    assert body["provider"] == "hosted" and body["input"]["media"] == "text" and len(body["input"]["sha256"]) == 64
    assert "content" not in body["input"] and "Harrowgate" not in json.dumps(body["input"])
    (batch,) = body["batches"]
    assert (batch["model_version"], batch["use_case"]) == ("MediCore-2.5", "Clinician Facing CDSS")
    assert [(r["record"]["code"], r["record"]["age"]) for r in batch["records"]] == [("5A11", "All ages"), ("CA22", "All ages"), ("5A11", "Age 65+")]
    first = batch["records"][0]
    assert first["fields"]["perf"] == {"origin": "study", "source": {"quote": "the accuracy was 0.83", "location": "Results, page 2"},
                                       "confidence": 0.95, "uncertain": False, "reason": None}
    # Fields the study does not state keep the stored row's values (a merge row replaces the whole row).
    assert first["fields"]["intended_use_level"]["origin"] == "store" and first["record"]["intended_use_level"] == "In scope"
    assert (first["record"]["source"], first["record"]["study"], first["record"]["source_date"]) == (
        "Harrowgate multisite evaluation of MediCore-2.5 as clinician-facing decision support (fictional)", "HMS-2026-07", "2026-10-02")
    assert first["conditions"] == [{"code": "5A11", "name": "Type 2 diabetes mellitus"}]

    e4 = propose_file(client, "E4").json()  # a code the atlas does not have
    assert [r["record"]["code"] for b in e4["batches"] for r in b["records"]] == ["8A80"]
    (dropped,) = e4["dropped"]
    assert dropped["reason"] == "unknown_code" and "ZZ99" in dropped["detail"] and dropped["values"]["code"] == "ZZ99"
    assert dropped["source"] == {"quote": "local code ZZ99", "location": "paragraph 2"}

    e5 = propose_file(client, "E5").json()  # vague wording stays flagged
    (record,) = e5["batches"][0]["records"]
    assert record["uncertain"] is True
    assert record["fields"]["perf"]["uncertain"] and "Approximate" in record["fields"]["perf"]["reason"]
    assert record["fields"]["sex"]["uncertain"]

    e6 = propose_file(client, "E6").json()  # an AI system that is not a card
    assert [d["reason"] for d in e6["dropped"]] == ["unknown_card"]
    (incident,) = e6["batches"][0]["records"]
    assert incident["record"]["incident_id"].startswith("AST-") and incident["fields"]["incident_id"]["origin"] == "generated"
    assert incident["record"]["codes"] == "5A11" and incident["record"]["status"] == "Mitigated"

    e7 = propose_file(client, "E7").json()  # names without codes: exact matches in the store
    assert [(r["record"]["code"], r["fields"]["code"]["origin"]) for r in e7["batches"][0]["records"]] == [("CA23", "search"), ("FA00", "search")]


def test_card_hint_fills_and_restricts(client):
    configure(client, provider="hosted", api_key=CANARY)
    use_fake(client)
    other = propose_file(client, "E1", model_version="SynGPT-2.0", use_case="Clinician Facing CDSS").json()
    assert other["batches"] == [] and {d["reason"] for d in other["dropped"]} == {"other_card"}
    unknown = propose_file(client, "E1", model_version="Nope-1", use_case="CDSS")
    assert unknown.status_code == 422 and unknown.json()["code"] == "invalid_parameter"
    half = propose_file(client, "E1", model_version="MediCore-2.5")
    assert half.status_code == 422
    # A record without a card takes the chosen one.
    answer = {"records": [{"code": {"value": "BA41", "quote": "heart attack", "location": "line 1"},
                           "perf": {"value": 0.9, "quote": "0.9", "location": "line 1"}}], "notes": []}
    text = "Synthetic: 0.9 for heart attack letters (fictional)."
    client.app.state.assist.provider_factory = lambda config: FakeProvider.for_study_set({from_text(text).sha256: answer})
    hinted = client.post(f"{ADMIN}/assist/proposals", data={"kind": "metrics", "text": text, "model_version": "Quillary-Scribe-1.0",
                                                           "use_case": "Discharge letter generation"}).json()
    (batch,) = hinted["batches"]
    assert batch["model_version"] == "Quillary-Scribe-1.0" and batch["records"][0]["record"]["code"] == "BA41"
    # Without Source, Study, and Source date the merge row is flagged (the check would refuse it).
    assert batch["records"][0]["fields"]["source"]["uncertain"]


def test_subcategories_and_chapters_are_not_guessed_into_conditions(client):
    configure(client, provider="hosted", api_key=CANARY)
    answer = {"records": [
        {"model_version": {"value": "SynGPT-2.0", "quote": "q", "location": "l"}, "use_case": {"value": "Clinician Facing CDSS", "quote": "q", "location": "l"},
         "code": {"value": "CA40.0", "quote": "bacterial pneumonia", "location": "l"}, "perf": {"value": "81%", "quote": "81%", "location": "l"}},
        {"model_version": {"value": "SynGPT-2.0", "quote": "q", "location": "l"}, "use_case": {"value": "Clinician Facing CDSS", "quote": "q", "location": "l"},
         "code": {"value": "5A11", "quote": "diabetes", "location": "l"}, "perf": {"value": "81%", "quote": "81%", "location": "l"},
         "age": {"value": "65+", "quote": "older", "location": "l"}},
        {"model_version": {"value": "SynGPT-2.0", "quote": "q", "location": "l"}, "use_case": {"value": "Clinician Facing CDSS", "quote": "q", "location": "l"},
         "code": {"value": "5A11", "quote": "diabetes", "location": "l"}, "age": {"value": "Age 120+", "quote": "q", "location": "l"}},
    ], "notes": []}
    text = "Synthetic subcategory study (fictional)."
    client.app.state.assist.provider_factory = lambda config: FakeProvider.for_study_set({from_text(text).sha256: answer})
    body = client.post(f"{ADMIN}/assist/proposals", data={"kind": "metrics", "text": text}).json()
    assert [d["reason"] for d in body["dropped"]] == ["not_a_condition", "invalid_value"]
    assert "CA40" in body["dropped"][0]["detail"]
    (record,) = body["batches"][0]["records"]
    assert record["record"]["age"] == "Age 65+" and record["record"]["perf"] == 0.81
    assert record["fields"]["perf"]["uncertain"] and "percentage" in record["fields"]["perf"]["reason"]


def fake_answer(client, text: str, answer: dict) -> None:
    client.app.state.assist.provider_factory = lambda config: FakeProvider.for_study_set({from_text(text).sha256: answer})


@pytest.mark.parametrize("bad", ["nan", "NaN", "inf", "-Infinity", float("nan"), float("inf"), 10**400, "nan%"])
def test_non_finite_numbers_become_flagged_empty_values(client, bad):
    """Review S5 #2: NaN and infinity (text or JSON numbers) never reach the encoder (500) or int() (502)."""
    configure(client, provider="hosted", api_key=CANARY)
    answer = copy.deepcopy(ANSWERS["E1"])
    answer["records"][0]["perf"]["value"] = bad
    answer["records"][0]["n_samples"]["value"] = bad
    answer["records"][0]["perf"]["confidence"] = float("nan")
    answer["records"][1]["perf"]["confidence"] = float("-inf")
    # The study is E1's own text, so its quotes are found (S8: a quote not in the study is flagged).
    text = f"{CASES['E1'].content().decode()}\nSynthetic non-finite probe {bad!r} (fictional)."
    fake_answer(client, text, answer)
    response = client.post(f"{ADMIN}/assist/proposals", data={"kind": "metrics", "text": text})
    assert response.status_code == 200, response.text  # the client validates every response against the contract
    assert "NaN" not in response.text and "Infinity" not in response.text
    first, second = response.json()["batches"][0]["records"][:2]
    assert first["fields"]["perf"]["confidence"] is None and second["fields"]["perf"]["confidence"] is None
    assert second["fields"]["perf"]["uncertain"] is False  # a dropped confidence alone is no flag
    for name in ("perf", "n_samples"):  # never filled from the stored row
        assert first["record"][name] is None and first["fields"][name]["origin"] == "study"
        assert first["fields"][name]["uncertain"] and "Not a number." in first["fields"][name]["reason"]


@pytest.mark.parametrize("bad", ["high", "nan", "about a third"])
def test_a_value_that_does_not_parse_stays_flagged_and_empty(client, bad):
    """Review S5 #3: an unparseable study value is never replaced by the stored one under the new study's source."""
    configure(client, provider="hosted", api_key=CANARY)
    answer = copy.deepcopy(ANSWERS["E1"])
    answer["records"][0]["perf"]["value"] = bad
    text = f"Synthetic unparseable probe {bad} (fictional)."
    fake_answer(client, text, answer)
    response = client.post(f"{ADMIN}/assist/proposals", data={"kind": "metrics", "text": text})
    assert response.status_code == 200, response.text
    record = response.json()["batches"][0]["records"][0]
    perf = record["fields"]["perf"]
    assert record["record"]["perf"] is None and record["uncertain"] is True
    assert perf["origin"] == "study" and perf["uncertain"] is True
    assert perf["reason"].startswith("Not a number.") and "The stored row has" in perf["reason"]
    assert perf["source"] == {"quote": "the accuracy was 0.83", "location": "Results, page 2"}
    # Fields the model did not give at all still come from the stored row.
    assert record["fields"]["intended_use_level"]["origin"] == "store"


def test_a_row_mixing_stored_values_with_this_studys_source_is_flagged(client):
    """Review S5 #9: staging credits the whole merge row to the new study, so the maintainer decides."""
    from app.assist.proposals import MIXED_SOURCES, STUDY_REPORTED_FIELDS

    configure(client, provider="hosted", api_key=CANARY)
    use_fake(client)
    body = propose_file(client, "E1").json()
    for record in body["batches"][0]["records"]:
        stored = [name for name in STUDY_REPORTED_FIELDS if record["fields"][name]["origin"] == "store"]
        assert stored, "E1 keeps study-reported values (CI bounds) from the stored row"
        source = record["fields"]["source"]
        assert source["origin"] == "study" and source["uncertain"] is True and record["uncertain"] is True
        assert source["reason"] == MIXED_SOURCES
        assert source["source"] is not None  # the quote stays
    # Without this study's Source, Study, or Source date nothing is credited to it: only "missing" is said.
    answer = copy.deepcopy(ANSWERS["E1"])
    for record in answer["records"]:
        for name in ("source", "study", "source_date"):
            record.pop(name, None)
    text = "Synthetic study without its own source (fictional)."
    fake_answer(client, text, answer)
    record = client.post(f"{ADMIN}/assist/proposals", data={"kind": "metrics", "text": text}).json()["batches"][0]["records"][0]
    assert record["fields"]["source"]["uncertain"] and MIXED_SOURCES not in record["fields"]["source"]["reason"]


def test_non_finite_values_in_left_out_items_are_empty(client):
    configure(client, provider="hosted", api_key=CANARY)
    answer = copy.deepcopy(ANSWERS["E4"])
    for record in answer["records"]:
        if record.get("code", {}).get("value") == "ZZ99":
            record["perf"]["value"] = float("nan")
    text = "Synthetic dropped non-finite probe (fictional)."
    fake_answer(client, text, answer)
    response = client.post(f"{ADMIN}/assist/proposals", data={"kind": "metrics", "text": text})
    assert response.status_code == 200, response.text
    (dropped,) = response.json()["dropped"]
    assert dropped["values"]["code"] == "ZZ99" and dropped["values"]["perf"] is None


# ------------------------------------------------- proposes, never publishes

def test_proposing_never_stages_or_publishes_and_accepting_stages_a_merge(client):
    configure(client, provider="hosted", api_key=CANARY)
    use_fake(client)
    versions = client.get(f"{ADMIN}/versions").json()
    before = client.get(f"{ADMIN}/datasets").json()
    proposals = propose_file(client, "E1").json()
    assert client.get(f"{ADMIN}/versions").json() == versions
    assert client.get(f"{ADMIN}/datasets").json() == before  # nothing staged

    (batch,) = proposals["batches"]
    accepted = client.post(f"{ADMIN}/uploads/records", json={
        "kind": "metrics", "model_version": batch["model_version"], "use_case": batch["use_case"], "mode": "merge",
        "records": [item["record"] for item in batch["records"]],
    })
    assert accepted.status_code == 202, accepted.text
    upload = wait(client, accepted.json()["upload_id"])
    assert upload["status"] == "valid", upload["report"]
    assert upload["mode"] == "merge" and upload["origin"] == "records"
    assert client.get(f"{ADMIN}/versions").json() == versions  # still nothing published
    changes = client.get(f"{ADMIN}/uploads/{upload['upload_id']}/changes").json()
    assert changes["rows"]["changed"] + changes["rows"]["added"] == 3
    published = client.post(f"{ADMIN}/uploads/{upload['upload_id']}/publish", json={"previewed": True, "note": "Harrowgate study (fictional)."})
    assert published.status_code == 200
    concept = client.get(f"{API_PREFIX}/metrics/concepts/CA22", params={"model": "MediCore-2.5", "use_case": "Clinician Facing CDSS"}).json()
    assert concept["metrics"]["provenance"]["data_source"]["study"] == "HMS-2026-07"

    # A safety proposal accepted the same way.
    incidents = propose_file(client, "E3").json()
    records = [item["record"] for item in incidents["batches"][0]["records"]]
    staged = client.post(f"{ADMIN}/uploads/records", json={"kind": "safety", "records": records})
    assert staged.status_code == 202
    assert wait(client, staged.json()["upload_id"])["status"] == "valid"
    client.delete(f"{ADMIN}/uploads/{staged.json()['upload_id']}")


# ------------------------------------------------------- providers over HTTP

def test_hosted_provider_runs_the_search_tool_loop_over_http(client, stub):
    configure(client, provider="hosted", api_key=CANARY, base_url=stub.url, model="claude-sonnet-5-5")
    stub.requests.clear()
    final = json.dumps(ANSWERS["E1"])

    def respond(path, headers, body):
        if headers.get("x-api-key") != CANARY or headers.get("anthropic-version") != "2023-06-01":
            return 401, {"type": "error", "error": {"type": "authentication_error", "message": "bad key"}}, {}
        if len(body["messages"]) == 1:
            return 200, anthropic_tool_use("toolu_1", "type 2 diabetes"), {}
        return 200, anthropic_text(final), {}

    stub.respond = respond
    response = propose_file(client, "E1")
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["requests"] == 2 and body["model"] == "claude-sonnet-5-5"
    assert len(body["batches"][0]["records"]) == 3
    first, second = stub.requests
    assert first["path"] == "/v1/messages" and first["body"]["model"] == "claude-sonnet-5-5"
    assert first["body"]["tools"][0]["name"] == "search_conditions" and "tool_choice" not in first["body"]
    assert "<study>" in first["body"]["messages"][0]["content"][0]["text"]
    assert "MediCore-2.5 | Clinician Facing CDSS" in first["body"]["system"]
    tool_result = second["body"]["messages"][2]["content"][0]
    assert tool_result["type"] == "tool_result" and tool_result["tool_use_id"] == "toolu_1"
    assert any(hit["code"] == "5A11" for hit in json.loads(tool_result["content"]))
    assert second["body"]["messages"][1]["role"] == "assistant"

    # A PDF goes as a base64 document block.
    stub.requests.clear()
    stub.respond = lambda path, headers, body: (200, anthropic_text('{"records": [], "notes": ["Nothing to propose."]}'), {})
    pdf = b"%PDF-1.4\n% synthetic test document\n"
    response = client.post(f"{ADMIN}/assist/proposals", data={"kind": "safety"}, files={"file": ("study.pdf", pdf, "application/pdf")})
    assert response.status_code == 200 and response.json()["input"]["media"] == "pdf"
    block = stub.requests[0]["body"]["messages"][0]["content"][0]
    assert block["type"] == "document" and block["source"]["media_type"] == "application/pdf"
    assert base64.b64decode(block["source"]["data"]) == pdf


def test_local_provider_resolves_codes_in_two_passes(client, stub):
    configure(client, provider="local", model="llama3.1:8b", base_url=f"{stub.url}/v1")
    stub.requests.clear()
    answer = json.loads(json.dumps(ANSWERS["E7"]))
    answer["records"][0]["condition"]["value"] = "asthma"  # an exact name, any case
    answer["records"][1]["condition"]["value"] = "osteoarthritis of the hip"  # not exact: the model chooses

    def respond(path, headers, body):
        if "authorization" in headers:
            return 400, {"error": {"message": "unexpected header"}}, {}
        prompt = body["messages"][-1]["content"]
        if "candidates" in prompt:
            mentions = json.loads(prompt[prompt.index("\n[") : prompt.index("\nAnswer with JSON only")])
            choices = [{"id": m["id"], "code": next((c["code"] for c in m["candidates"] if c["name"] == "Osteoarthritis of hip"), None)} for m in mentions]
            return 200, openai_text(json.dumps({"choices": choices})), {}
        return 200, openai_text("```json\n" + json.dumps(answer) + "\n```"), {}

    stub.respond = respond
    response = propose_file(client, "E7")
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["requests"] == 2 and body["provider"] == "local"
    records = body["batches"][0]["records"]
    assert [r["record"]["code"] for r in records] == ["CA23", "FA00"]
    assert records[0]["fields"]["code"]["uncertain"] is False
    assert records[1]["fields"]["code"]["uncertain"] is True  # chosen by the model among store candidates
    assert stub.requests[0]["path"] == "/v1/chat/completions" and stub.requests[0]["body"]["temperature"] == 0
    assert "search_conditions" not in json.dumps(stub.requests[0]["body"])

    # A PDF never reaches a local model.
    stub.requests.clear()
    pdf = client.post(f"{ADMIN}/assist/proposals", data={"kind": "metrics"}, files={"file": ("study.pdf", b"%PDF-1.4 x", "application/pdf")})
    assert pdf.status_code == 422 and pdf.json()["code"] == "assist_unsupported_input"
    assert "paste the text" in pdf.json()["detail"] and stub.requests == []


@pytest.mark.parametrize(
    ("status", "answer", "headers", "fragment"),
    [
        (401, {"type": "error", "error": {"type": "authentication_error", "message": f"invalid x-api-key {CANARY}"}}, {}, "rejected"),
        (429, {"type": "error", "error": {"type": "rate_limit_error", "message": "slow down"}}, {}, "rate limiting"),
        (500, {"type": "error", "error": {"type": "api_error", "message": f"echo {CANARY}"}}, {}, "HTTP 500"),
        (302, {}, {"Location": "http://127.0.0.1:9/steal"}, "redirect"),
        (200, anthropic_text("I cannot find any records."), {}, "no JSON"),
        (200, {**anthropic_text(""), "stop_reason": "refusal"}, {}, "declined"),
        (200, {**anthropic_text("{"), "stop_reason": "max_tokens"}, {}, "cut off"),
        # Review fix: an answer nested too deeply to read is a provider problem, not a server error.
        (200, anthropic_text('{"a": ' + "[" * 300_000 + "]" * 300_000 + "}"), {}, "nested too deeply"),
    ],
)
def test_provider_errors_are_problems_without_the_key(client, stub, status, answer, headers, fragment):
    configure(client, provider="hosted", api_key=CANARY, base_url=stub.url)
    stub.requests.clear()
    stub.respond = lambda path, request_headers, body: (status, answer, headers)
    response = propose_file(client, "E1")
    assert response.status_code == 502, response.text
    assert response.json()["code"] == "assist_provider_error" and fragment in response.json()["detail"]
    assert CANARY not in response.text
    assert [request["path"] for request in stub.requests] == ["/v1/messages"]  # a redirect is not followed


def test_logs_carry_the_helpers_words_never_the_providers_or_the_studys(client, stub, logs):
    """Review S5 #7: provider messages and unexpected error texts stay out of the log."""
    configure(client, provider="hosted", api_key=CANARY, base_url=stub.url)
    marker = "provider-says-" + secrets.token_hex(6)
    stub.respond = lambda path, headers, body: (500, {"type": "error", "error": {"type": "api_error", "message": f"{marker} {CANARY}"}}, {})
    logs.lines.clear()
    response = propose_file(client, "E1")
    assert response.status_code == 502 and marker in response.json()["detail"] and CANARY not in response.text
    assert any("The provider answered HTTP 500." in line for line in logs.lines)
    assert not [line for line in logs.lines if marker in line or CANARY in line]

    study_words = "study-words-" + secrets.token_hex(6)

    class Exploding(FakeProvider):
        def extract(self, study, context, tools):
            raise RuntimeError(f"{study_words} {CANARY}")

    client.app.state.assist.provider_factory = lambda config: Exploding(kind="hosted", model="boom", accepts_pdf=True, uses_search_tool=True)
    logs.lines.clear()
    response = client.post(f"{ADMIN}/assist/proposals", data={"kind": "metrics", "text": "Synthetic."})
    assert response.status_code == 502 and study_words not in response.text
    assert any("Assisted entry failed: RuntimeError." in line for line in logs.lines)
    assert not [line for line in logs.lines if study_words in line or CANARY in line]


def test_one_proposal_at_a_time(client):
    """Review S5 #8: a cancelled request keeps running, so a retry is refused (409 assist_busy) instead of piling up."""
    import threading

    configure(client, provider="hosted", api_key=CANARY)
    entered, release = threading.Event(), threading.Event()

    class Slow(FakeProvider):
        def extract(self, study, context, tools):
            entered.set()
            release.wait(10)
            return super().extract(study, context, tools)

    client.app.state.assist.provider_factory = lambda config: Slow(kind="hosted", model="slow", accepts_pdf=True, uses_search_tool=True)
    service = client.app.state.assist
    first = {}
    worker = threading.Thread(target=lambda: first.setdefault("result", service.propose(kind="metrics", study=from_text("Synthetic slow study."))))
    worker.start()
    try:
        assert entered.wait(10)
        busy = client.post(f"{ADMIN}/assist/proposals", data={"kind": "metrics", "text": "Synthetic second study."})
        assert busy.status_code == 409 and busy.json()["code"] == "assist_busy", busy.text
        assert busy.json()["title"] == "The helper is reading another study" and "one at a time" in busy.json()["detail"]
        assert client.get(f"{ADMIN}/assist").status_code == 200  # status and settings stay available
    finally:
        release.set()
        worker.join(10)
    assert first["result"]["model"] == "slow"
    # Free again once the first one is done, also after a failure.
    assert client.post(f"{ADMIN}/assist/proposals", data={"kind": "metrics", "text": "Synthetic third study."}).status_code == 200


def test_a_proposal_has_one_overall_deadline(client, stub, monkeypatch):
    """Review S5 #8: the tool loop and the chooser share one wall-clock budget."""
    import time

    from app.assist import providers

    monkeypatch.setattr(providers, "TOTAL_SECONDS", 0.6)
    monkeypatch.setattr(providers, "TIMEOUT_SECONDS", 0.6)
    configure(client, provider="hosted", api_key=CANARY, base_url=stub.url)
    stub.requests.clear()

    def slow_search(path, headers, body):
        time.sleep(0.25)
        return 200, anthropic_tool_use(f"toolu_{len(body['messages'])}", "asthma"), {}

    stub.respond = slow_search
    started = time.perf_counter()
    response = propose_file(client, "E1")
    assert response.status_code == 502 and response.json()["code"] == "assist_provider_error", response.text
    assert "the most one proposal may take" in response.json()["detail"]
    assert 2 <= len(stub.requests) < providers.MAX_TOOL_TURNS and time.perf_counter() - started < 3

    # A request that outlasts what is left is cut at the deadline, with the same words.
    configure(client, provider="local", model="m", base_url=f"{stub.url}/v1")
    stub.respond = lambda path, headers, body: (time.sleep(1.5), (200, openai_text('{"records": [], "notes": []}'), {}))[1]
    started = time.perf_counter()
    response = client.post(f"{ADMIN}/assist/proposals", data={"kind": "metrics", "text": "Synthetic slow local study."})
    assert response.status_code == 502 and "the most one proposal may take" in response.json()["detail"], response.text
    assert time.perf_counter() - started < 1.5


def test_unreachable_provider_is_a_problem(client):
    configure(client, provider="local", model="m", base_url="http://127.0.0.1:9/v1", api_key=CANARY)
    response = client.post(f"{ADMIN}/assist/proposals", data={"kind": "metrics", "text": "Synthetic text."})
    assert response.status_code == 502 and "reach" in response.json()["detail"] and CANARY not in response.text


# ---------------------------------------------------------------- inputs

def test_inputs_tables_and_limits(client):
    csv_input = from_file("table.csv", CASES["E2"].content())
    assert csv_input.media == "csv" and "row 4: Acute myocardial infarction | BA41" in csv_input.text
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.active.title = "Results"
    workbook.active.append(["Condition", "Share"])
    workbook.active.append(["Pneumonia", 0.93])
    buffer = io.BytesIO()
    workbook.save(buffer)
    xlsx = from_file("table.xlsx", buffer.getvalue())
    assert xlsx.media == "xlsx" and "Sheet Results" in xlsx.text and "row 2: Pneumonia | 0.93" in xlsx.text

    configure(client, provider="hosted", api_key=CANARY)
    use_fake(client)
    assert client.post(f"{ADMIN}/assist/proposals", data={"kind": "metrics", "text": "  "}).status_code == 422
    both = client.post(f"{ADMIN}/assist/proposals", data={"kind": "metrics", "text": "x"}, files={"file": ("a.txt", b"x", "text/plain")})
    assert both.status_code == 422
    assert client.post(f"{ADMIN}/assist/proposals", data={"kind": "models", "text": "x"}).status_code == 422
    fake_pdf = client.post(f"{ADMIN}/assist/proposals", data={"kind": "metrics"}, files={"file": ("a.pdf", b"not a pdf", "application/pdf")})
    assert fake_pdf.status_code == 415
    big = client.post(f"{ADMIN}/assist/proposals", data={"kind": "metrics"}, files={"file": ("big.txt", b"a" * (11 * 1024 * 1024), "text/plain")})
    assert big.status_code == 413 and big.json()["code"] == "upload_too_large"
    pasted = client.post(f"{ADMIN}/assist/proposals", data={"kind": "metrics", "text": CASES["E1"].content().decode()})
    assert pasted.status_code == 200 and len(pasted.json()["batches"][0]["records"]) == 3  # the same study, pasted


def _xlsx(build) -> bytes:
    from openpyxl import Workbook

    workbook = Workbook()
    build(workbook)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_unreadable_tables_are_a_plain_422(client):
    """Review S5 #4: csv and openpyxl errors are input problems, not 500s."""
    configure(client, provider="hosted", api_key=CANARY)
    use_fake(client)
    big_cell = b"a,b\n" + b"x" * 200_000 + b",1\n"  # longer than the csv module's field limit
    response = client.post(f"{ADMIN}/assist/proposals", data={"kind": "metrics"}, files={"file": ("t.csv", big_cell, "text/csv")})
    assert response.status_code == 422 and response.json()["code"] == "assist_unsupported_input", response.text
    assert "could not be read as CSV" in response.json()["detail"] and "xxxx" not in response.text
    late_nul = b"a,b\n" + b"y,1\n" * 2000 + b"z\x00,2\n"
    response = client.post(f"{ADMIN}/assist/proposals", data={"kind": "metrics"}, files={"file": ("t.csv", late_nul, "text/csv")})
    assert response.status_code in (200, 422), response.text
    # Workbooks that open but whose rows cannot be read (IndexError, ValueError while streaming), and one that does not open.
    import zipfile

    good = _xlsx(lambda workbook: workbook.active.append(["Condition", "Share"]))
    sheet = '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData><row r="1"><c r="A1" t="{t}"><v>{v}</v></c></row></sheetData></worksheet>'
    broken = []
    for cell in ({"t": "s", "v": "99"}, {"t": "n", "v": "abc"}):
        buffer = io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(good)) as source, zipfile.ZipFile(buffer, "w") as target:
            for item in source.infolist():
                target.writestr(item, sheet.format(**cell).encode() if item.filename.endswith("sheet1.xml") else source.read(item))
        broken.append(buffer.getvalue())
    for content in (*broken, b"PK\x03\x04 not really a workbook"):
        response = client.post(f"{ADMIN}/assist/proposals", data={"kind": "metrics"}, files={"file": ("t.xlsx", content, "application/octet-stream")})
        assert response.status_code == 422 and response.json()["code"] == "assist_unsupported_input", response.text
        assert response.json()["detail"] == "The workbook could not be read."


def test_large_tables_are_streamed_and_cut_and_the_response_says_so(client):
    """Review S5 #5: rows, columns, sheets, and characters are capped; the input says it was cut."""
    from app.assist import inputs

    rows = from_file("rows.csv", ("a,b\n" + "".join(f"c{n},{n}\n" for n in range(2, 2600))).encode())
    assert rows.truncated and "row 2000: c2000 | 2000" in rows.text and "row 2001" not in rows.text
    assert "only the first 2,000 rows" in rows.text
    small = from_file("small.csv", b"a,b\nc,1\n")
    assert small.truncated is False and small.info()["truncated"] is False

    wide = from_file("wide.xlsx", _xlsx(lambda workbook: workbook.active.append([f"h{n}" for n in range(150)])))
    assert wide.truncated and "h99" in wide.text and "h100" not in wide.text and "only the first 100" in wide.text

    def many_sheets(workbook):
        for number in range(25):
            sheet = workbook.active if number == 0 else workbook.create_sheet()
            sheet.title = f"S{number}"
            sheet.append([f"value {number}"])

    sheets = from_file("sheets.xlsx", _xlsx(many_sheets))
    assert sheets.truncated and "Sheet S19" in sheets.text and "Sheet S20" not in sheets.text and "only the first 20" in sheets.text

    # A sparse sheet with a far-away row is read only up to the row limit, quickly.
    import time

    def far(workbook):
        workbook.active.append(["Condition", "Share"])
        workbook.active.cell(row=1_000_000, column=1, value="far away")

    started = time.perf_counter()
    sparse = from_file("sparse.xlsx", _xlsx(far))
    assert time.perf_counter() - started < 5 and sparse.truncated and "far away" not in sparse.text

    # File-derived text has the same cap as pasted text.
    long_text = from_file("long.txt", ("Synthetic study sentence (fictional). " * 7000).encode())
    assert long_text.truncated and len(long_text.text) <= inputs.MAX_TEXT_CHARS + len(inputs.CUT_NOTE) + 1
    assert long_text.text.endswith(inputs.CUT_NOTE)
    long_table = from_file("long.csv", ("a,b\n" + "".join(f"{'w' * 150},{n}\n" for n in range(1900))).encode())
    assert long_table.truncated and long_table.text.endswith(inputs.CUT_NOTE)

    configure(client, provider="hosted", api_key=CANARY)
    use_fake(client)
    content = ("Synthetic study sentence (fictional). " * 7000).encode()
    response = client.post(f"{ADMIN}/assist/proposals", data={"kind": "metrics"}, files={"file": ("long.txt", content, "text/plain")})
    assert response.status_code == 200, response.text
    assert response.json()["input"]["truncated"] is True and response.json()["input"]["bytes"] == len(content)
    whole = propose_file(client, "E1").json()
    assert whole["input"]["truncated"] is False


def test_non_loopback_clients_are_refused(demo_store):
    from fastapi.testclient import TestClient

    with TestClient(create_app(demo_store), base_url="http://127.0.0.1", client=("203.0.113.9", 4000)) as remote:
        assert remote.get(f"{ADMIN}/assist").status_code == 403
        assert remote.put(f"{ADMIN}/assist", json={"provider": "hosted", "api_key": CANARY}).status_code == 403


# ----------------------------------------------------------- evaluation

def test_fake_evaluation_reports_field_accuracy(demo_store, capsys):
    assert evaluate.main(["--db", str(demo_store), "--provider", "fake", "--json"]) == 0
    result = json.loads(capsys.readouterr().out)
    summary = result["summary"]
    assert summary["cases"] == len(CASES) == 7 and summary["errors"] == 0
    assert summary["fields_total"] == 194 and summary["fields_correct"] == 193  # E5's ambiguous group is the miss
    assert summary["record_recall"] == 1 and summary["record_precision"] == 1
    assert (summary["drops_found"], summary["drops_expected"]) == (2, 2)
    assert (summary["flags_found"], summary["flags_expected"]) == (2, 2)
    assert summary["codes_in_store"] == summary["codes_proposed"] > 0
    assert summary["sourced_fields"] == summary["study_fields"]
    assert evaluate.main(["--db", str(demo_store), "--provider", "fake", "--cases", "E1,E4"]) == 0
    assert "Field accuracy:      100.0% (67/67)" in capsys.readouterr().out


def test_real_provider_evaluation_needs_a_configuration(demo_store, tmp_path, capsys):
    assert evaluate.main(["--db", str(demo_store), "--provider", "hosted", "--config", str(tmp_path / "none.json")]) == 2
    assert "No hosted helper configuration" in capsys.readouterr().err


def test_study_set_is_synthetic_and_fake_answers_cover_it():
    keyed = fake_answers()
    for case in CASES.values():
        text = case.content().decode()
        assert "SYNTHETIC" in text and ("fictional" in text.lower())
        assert case.study().sha256 in keyed
    assert {case.kind for case in CASES.values()} == {"metrics", "safety"}
    assert any(case.path.suffix == ".csv" for case in CASES.values())


# ------------------------------------------------------ the key never leaks

def test_the_key_never_reaches_logs_store_or_uploads(client, demo_store, logs, stub):
    # Exercise every path once more with the canary key, then look everywhere.
    configure(client, provider="hosted", api_key=CANARY, base_url=stub.url)
    stub.respond = lambda path, headers, body: (401, {"type": "error", "error": {"type": "authentication_error", "message": CANARY}}, {})
    assert propose_file(client, "E3").status_code == 502
    stub.respond = lambda path, headers, body: (200, anthropic_text(json.dumps(ANSWERS["E3"])), {})
    proposals = propose_file(client, "E3").json()
    staged = client.post(f"{ADMIN}/uploads/records", json={"kind": "safety", "records": [r["record"] for r in proposals["batches"][0]["records"]]})
    upload = wait(client, staged.json()["upload_id"])
    for response in (
        client.get(f"{ADMIN}/assist"),
        client.get(f"{ADMIN}/uploads/{upload['upload_id']}"),
        client.get(f"{ADMIN}/uploads/{upload['upload_id']}/annotated.csv"),
        client.get(f"{ADMIN}/templates/safety.csv", params={"filled": "true"}),
        client.get(f"{API_PREFIX}/export/concepts/BA41", params={"model": "Vantrel-Triage-2.1", "use_case": "Emergency department triage"}),
        client.get(f"{API_PREFIX}/meta"),
    ):
        assert response.status_code == 200 and CANARY not in response.text
    client.delete(f"{ADMIN}/uploads/{upload['upload_id']}")

    assert logs.lines, "the module's log records were collected"
    assert not [line for line in logs.lines if CANARY in line]
    # Requests to the provider carried it (the stub saw it), nothing else did.
    assert any(request["headers"].get("x-api-key") == CANARY for request in stub.requests)
    store_dir = Path(demo_store).parent
    for path in store_dir.rglob("*"):
        if path.is_file() and path != client.app.state.assist.config_path:
            assert CANARY.encode() not in path.read_bytes(), f"the key leaked into {path}"
    uploads = client.app.state.api.admin.uploads_dir
    assert uploads.exists() and any(uploads.rglob("*"))


def test_a_quote_that_is_not_in_the_study_is_flagged(client):
    """S8 critique: a quote shown as evidence must be words of the study as read; otherwise the field is flagged."""
    configure(client, provider="hosted", api_key=CANARY)
    study = CASES["E1"].content().decode()
    answer = copy.deepcopy(ANSWERS["E1"])
    answer["records"][0]["perf"]["quote"] = "accuracy was 0.99 in a cohort the study never had"
    fake_answer(client, study, answer)
    response = client.post(f"{ADMIN}/assist/proposals", data={"kind": "metrics", "text": study})
    assert response.status_code == 200, response.text
    first, second = response.json()["batches"][0]["records"][:2]
    assert first["fields"]["perf"]["uncertain"] and "not in the study as read" in first["fields"]["perf"]["reason"]
    # Quotes that are the study's words (any case, spacing, or punctuation) are not flagged for it.
    assert "not in the study as read" not in (second["fields"]["perf"]["reason"] or "")


def test_quote_matching_ignores_case_spacing_and_punctuation():
    from app.assist.proposals import _words, quote_in

    study = _words("Sensitivity was 0.81 (95% CI 0.76-0.86) across 1,240 cases.")
    assert quote_in("sensitivity was 0.81", study)
    assert quote_in("95% CI 0.76–0.86", study)
    assert quote_in("1240 cases", study) is False  # "1,240" reads as "1 240": digits are compared as written
    assert quote_in("1,240 cases", study)
    assert not quote_in("accuracy was 0.83", study)
