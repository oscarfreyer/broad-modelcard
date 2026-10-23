"""Registry entry fields (contract v1.5).

Models and products with kind, AI type, developer, release date, description,
and built on; per configuration a function category and developer
declarations. All optional: registry files without the new columns stay
valid, a filled template round-trips without changes, and every new rule
rejects with row, column, and a plain-language fix. Runs on private copies
of the tiny store.
"""

from __future__ import annotations

import io
import sqlite3

import openpyxl
import pytest
from admin_helpers import (
    ADMIN,
    TINY_TARGET,
    errors_by_code,
    read_csv,
    template,
    upload_and_wait,
    write_csv,
)
from test_admin_rules import FIRST_DATA_ROW, check_marked, check_rejected

from app import db, repository
from app.main import API_PREFIX
from app.schema.templates import TEMPLATES

# The eight columns of contract v1.4. "Use case" and "Default use case" became "Function" and
# "Default function" in v1.7; the old headers stay accepted.
OLD_HEADER = ["Model version", "Display name", "Model name", "Release", "Status", "Use case", "Default use case", "Default model"]
CURRENT_HEADER = ["Model version", "Display name", "Model name", "Release", "Status", "Function", "Default function", "Default model"]
NEW_HEADERS = [
    "Kind", "AI type", "Developer", "Release date", "Description", "Built on", "Function category",
    "Declared intended purpose", "Declared regulatory status", "Declared contact", "Declared documentation URL",
    "Intended use stated",  # contract v1.11
    # Contract v1.14 (§28): lineage, the configuration record, the measure, and data protection.
    "Replaces", "Configuration ID", "Configuration date", "Prompt summary", "Prompt SHA-256", "Retrieval", "Tools",
    "Sampling temperature", "Sampling top-p", "Max output tokens", "Fine-tuning", "Performance measure",
    "Reference standard", "Declared data use", "Declared data location", "Declared retention", "Declared privacy contact",
]

ENTRY = {
    "Kind": "Model",
    "AI type": "Large language model",
    "Developer": "Tinyworks Labs (fictional)",
    "Release date": "2025-11-03",
    "Description": "A tiny fictional model for tests.",
    "Function category": "Decision support",
    "Declared intended purpose": "Fictional decision support for tests.",
    "Declared regulatory status": "Not assessed (demonstration)",
    "Declared contact": "tests@tinyworks.example",
    "Declared documentation URL": "https://tinyworks.example/docs",
}


def models(client) -> dict:
    response = client.get(f"{API_PREFIX}/models")
    assert response.status_code == 200, response.text
    return response.json()


def registry_csv(client, edit=None) -> bytes:
    header, rows, comments = read_csv(template(client, "models", filled="true"))
    if edit:
        edit(header, rows)
    return write_csv(header, rows, comments)


def counts(changes: dict) -> dict:
    return {key: value for key, value in changes["rows"].items() if key != "samples"}


def set_fields(values: dict, row: int = 0):
    def edit(header, rows):
        for column, value in values.items():
            rows[row][header.index(column)] = value

    return edit


def publish(client, upload: dict) -> dict:
    assert upload["status"] == "valid", upload["report"]
    response = client.post(f"{ADMIN}/uploads/{upload['upload_id']}/publish", json={"previewed": False})
    assert response.status_code == 200, response.text
    return response.json()


# ------------------------------------------------------------ compatibility

def test_seed_entries_are_models_without_the_new_fields(admin_client):
    model = models(admin_client)["models"][0]
    assert model["kind"] == "model"
    assert model["built_on"] == []
    assert all(model[field] is None for field in ("ai_type", "developer", "release_date", "description"))
    assert model["use_cases"][0]["function_category"] is None
    assert model["use_cases"][0]["declarations"] is None


def test_a_registry_file_without_the_new_columns_still_validates(admin_client):
    """An old-format file (the eight columns of contract v1.4) is valid and changes nothing."""
    content = write_csv(OLD_HEADER, [["Tiny-1", "Tiny-1", "Tiny", "1", "ACTIVE", "CDSS", "Yes", "Yes"]])
    upload = upload_and_wait(admin_client, "models", content)
    assert upload["status"] == "valid", upload["report"]
    assert upload["report"]["errors"] == [] and upload["report"]["notes"] == []
    changes = admin_client.get(f"{ADMIN}/uploads/{upload['upload_id']}/changes").json()
    assert counts(changes) == {"added": 0, "changed": 0, "removed": 0, "unchanged": 1}
    assert changes["models"]["changed"] == []


def test_the_template_has_the_new_columns_after_the_old_ones(admin_client):
    assert [column.header for column in TEMPLATES["models"]] == CURRENT_HEADER + NEW_HEADERS
    assert all(not column.required for column in TEMPLATES["models"][len(CURRENT_HEADER):])
    header, rows, _ = read_csv(template(admin_client, "models", filled="true"))
    assert header == CURRENT_HEADER + NEW_HEADERS
    assert rows[0][header.index("Kind")] == "Model"  # an entry without a kind is a model

    workbook = openpyxl.load_workbook(io.BytesIO(template(admin_client, "models", "xlsx")))
    instructions = [row[0] for row in workbook["Instructions"].iter_rows(values_only=True) if row and row[0]]
    assert "Built on" in instructions and "Declared documentation URL" in instructions
    lists = {value for row in workbook["Lists"].iter_rows(values_only=True) for value in row if value}
    assert {"Product", "Large language model", "Tabular risk model", "Triage", "Image analysis"} <= lists


# ------------------------------------------------------------ valid uploads

def test_entry_and_configuration_fields_publish_and_are_served(admin_client):
    client = admin_client
    upload = upload_and_wait(client, "models", registry_csv(client, set_fields(ENTRY)))
    changes = client.get(f"{ADMIN}/uploads/{upload['upload_id']}/changes").json()
    assert changes["rows"]["changed"] == 1 and changes["models"]["changed"] == ["Tiny-1"]
    publish(client, upload)

    model = models(client)["models"][0]
    assert {key: model[key] for key in ("kind", "ai_type", "developer", "release_date", "description", "built_on")} == {
        "kind": "model",
        "ai_type": "llm",
        "developer": "Tinyworks Labs (fictional)",
        "release_date": "2025-11-03",
        "description": "A tiny fictional model for tests.",
        "built_on": [],
    }
    use_case = model["use_cases"][0]
    assert use_case["function_category"] == "decision_support"
    assert use_case["declarations"] == {
        "declared_intended_purpose": "Fictional decision support for tests.",
        "declared_regulatory_status": "Not assessed (demonstration)",
        "declared_contact": "tests@tinyworks.example",
        "declared_documentation_url": "https://tinyworks.example/docs",
    }

    # The filled template carries the fields, and uploading it unchanged changes nothing, in both formats.
    for file_format in ("csv", "xlsx"):
        filled = template(client, "models", file_format, filled="true")
        again = upload_and_wait(client, "models", filled, f"registry.{file_format}", replace_staged=True)
        assert again["status"] == "valid", again["report"]
        summary = client.get(f"{ADMIN}/uploads/{again['upload_id']}/changes").json()
        assert counts(summary) == {"added": 0, "changed": 0, "removed": 0, "unchanged": 1}, file_format
        assert summary["models"]["changed"] == []


@pytest.mark.parametrize(
    ("column", "value", "field", "expected"),
    [
        ("Function category", "Triage", "function_category", "triage"),
        ("Declared contact", "someone@tinyworks.example", "declarations", "someone@tinyworks.example"),
        ("Developer", "Another Lab (fictional)", "developer", "Another Lab (fictional)"),
    ],
)
def test_changing_one_new_field_changes_the_configuration(admin_client, column, value, field, expected):
    client = admin_client
    publish(client, upload_and_wait(client, "models", registry_csv(client, set_fields(ENTRY))))
    upload = upload_and_wait(client, "models", registry_csv(client, set_fields({column: value})))
    changes = client.get(f"{ADMIN}/uploads/{upload['upload_id']}/changes").json()
    assert changes["rows"]["changed"] == 1 and changes["rows"]["samples"]["changed"][0]["name"] == "Tiny-1 / CDSS"
    publish(client, upload)
    model = models(client)["models"][0]
    served = model["use_cases"][0][field] if field in ("function_category", "declarations") else model[field]
    if field == "declarations":
        served = served["declared_contact"]
    assert served == expected


def test_codes_and_readable_values_are_both_accepted(admin_client):
    edit = set_fields({"Kind": "model", "AI type": "rule_based", "Function category": "image_analysis", "Release date": "2025-01-31"})
    upload = upload_and_wait(admin_client, "models", registry_csv(admin_client, edit))
    assert upload["status"] == "valid", upload["report"]


def test_a_date_cell_in_a_workbook_is_accepted(admin_client):
    from datetime import datetime

    header, rows, _ = read_csv(template(admin_client, "models", filled="true"))
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(header)
    row = list(rows[0])
    row[header.index("Release date")] = datetime(2025, 6, 1)
    sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    upload = upload_and_wait(admin_client, "models", buffer.getvalue(), "registry.xlsx")
    assert upload["status"] == "valid", upload["report"]
    publish(admin_client, upload)
    assert models(admin_client)["models"][0]["release_date"] == "2025-06-01"


def test_a_product_is_built_on_a_model_and_added_metrics_first(admin_client):
    """A product (kind Product, built on Tiny-1) joins the registry the same way a new model does."""
    client = admin_client

    def add_product(header, rows):
        product = [None] * len(header)
        values = {
            "Model version": "TinyCare-1", "Display name": "TinyCare 1", "Status": "EXPERIMENTAL", "Function": "Triage",
            "Default function": "Yes", "Default model": "No", "Kind": "Product", "AI type": "Hybrid",
            "Developer": "Tinyworks Labs (fictional)", "Built on": "Tiny-1", "Function category": "Triage",
            "Declared regulatory status": "Not assessed (demonstration)",
        }
        for column, value in values.items():
            product[header.index(column)] = value
        rows.append(product)

    registry = upload_and_wait(client, "models", registry_csv(client, add_product))
    assert list(errors_by_code(registry)) == ["use_case_without_metrics"]
    metrics = upload_and_wait(
        client, "metrics", template(client, "metrics", filled="true", **TINY_TARGET), model_version="TinyCare-1", use_case="Triage"
    )
    publish(client, metrics)
    client.app.state.api.admin.wait_for_checks()
    registry = client.get(f"{ADMIN}/uploads/{registry['upload_id']}").json()
    changes = client.get(f"{ADMIN}/uploads/{registry['upload_id']}/changes").json()
    assert changes["models"]["added"] == ["TinyCare-1"]
    publish(client, registry)

    product = models(client)["models"][1]
    assert product["model_version"] == "TinyCare-1"
    assert product["kind"] == "product" and product["built_on"] == ["Tiny-1"] and product["ai_type"] == "hybrid"
    assert product["use_cases"] == [
        {
            "use_case": "Triage",
            "has_multimorbidity": False,
            "function_category": "triage",
            "declarations": {
                "declared_intended_purpose": None,
                "declared_regulatory_status": "Not assessed (demonstration)",
                "declared_contact": None,
                "declared_documentation_url": None,
            },
            "intended_use_defined": True,  # contract v1.11
        }
    ]
    # The cdss-governance safeguard stays keyed on the use-case string, not on the kind.
    concept = client.get(f"{API_PREFIX}/metrics/concepts/1A00", params={"model": "TinyCare-1", "use_case": "Triage"}).json()
    assert "cdss-governance" not in [safeguard["id"] for safeguard in concept["metrics"]["safeguards"]]


# ----------------------------------------------------------- invalid uploads

CELL_CASES = [
    ("unknown_kind", "Kind", "Service", "Model or Product"),
    ("unknown_ai_type", "AI type", "Quantum model", "Large language model"),
    ("unknown_function_category", "Function category", "Billing", "Decision support"),
    ("invalid_date", "Release date", "2026-02-30", "YYYY-MM-DD"),
    ("invalid_date", "Release date", "03/01/2026", "YYYY-MM-DD"),
    ("invalid_url", "Declared documentation URL", "ftp://tinyworks.example/docs", "https://"),
    ("invalid_url", "Declared documentation URL", "www.tinyworks.example", "https://"),
    ("invalid_url", "Declared documentation URL", "http://[docs.example/x", "https://"),
    ("invalid_value", "Description", "x" * 281, "longer than 280"),
    ("invalid_value", "Declared intended purpose", "x" * 501, "longer than 500"),
    ("built_on_for_model", "Built on", "Tiny-1", "is a model"),
]


@pytest.mark.parametrize(("code", "column", "value", "problem"), CELL_CASES, ids=[f"{case[0]}-{case[2][:12]}" for case in CELL_CASES])
def test_invalid_entry_cells_are_rejected(admin_client, code, column, value, problem):
    upload = upload_and_wait(admin_client, "models", registry_csv(admin_client, set_fields({column: value})))
    check_rejected(admin_client, upload, code, FIRST_DATA_ROW, column, problem)
    check_marked(admin_client, upload["upload_id"], FIRST_DATA_ROW, column)
    assert errors_by_code(upload)[code]["fix"]


@pytest.mark.parametrize(
    ("built_on", "problem"),
    [("TinyCare-1", "built on itself"), ("Nope-9", "not a model version in this file"), ("Other-P", "is a product")],
)
def test_a_product_must_be_built_on_other_models_of_the_file(admin_client, built_on, problem):
    def edit(header, rows):
        for name in ("TinyCare-1", "Other-P"):
            product = list(rows[0])
            product[header.index("Model version")] = name
            product[header.index("Display name")] = name
            product[header.index("Default model")] = "No"
            product[header.index("Kind")] = "Product"
            product[header.index("Built on")] = built_on if name == "TinyCare-1" else "Tiny-1"
            rows.append(product)

    upload = upload_and_wait(admin_client, "models", registry_csv(admin_client, edit))
    group = errors_by_code(upload)["unknown_built_on"]
    assert group["examples"][0]["row"] == FIRST_DATA_ROW + 1 and group["examples"][0]["column"] == "Built on"
    assert problem in group["examples"][0]["problem"]


def test_entry_fields_must_agree_on_every_row_of_an_entry(admin_client):
    """Entry fields are repeated per row; a conflicting (or missing) repeat is a blocking error."""

    def edit(header, rows):
        rows[0][header.index("Developer")] = "Tinyworks Labs (fictional)"
        second = list(rows[0])
        second[header.index("Function")] = "Triage"
        second[header.index("Default function")] = "No"
        second[header.index("Developer")] = "Elsewhere Labs (fictional)"
        rows.append(second)
        third = list(second)
        third[header.index("Function")] = "Letters"
        third[header.index("Developer")] = ""
        rows.append(third)

    upload = upload_and_wait(admin_client, "models", registry_csv(admin_client, edit))
    group = errors_by_code(upload)["inconsistent_value"]
    assert [example["row"] for example in group["examples"]] == [FIRST_DATA_ROW + 1, FIRST_DATA_ROW + 2]
    assert all(example["column"] == "Developer" for example in group["examples"])
    assert "differs from row 4" in group["examples"][0]["problem"]
    assert "every row" in group["fix"] and "empty cell" in group["fix"]


# ---------------------------------------------------------------- migration

def test_a_version_2_store_gains_the_entry_columns_and_reads_as_before(tmp_path):
    """Migration 0003 is additive: existing registry rows read as models without the new fields."""
    conn = sqlite3.connect(tmp_path / "v2.sqlite", isolation_level=None)
    conn.execute("CREATE TABLE schema_migrations (version INTEGER PRIMARY KEY, name TEXT NOT NULL, applied_at TEXT NOT NULL)")
    for version, name, path in db.available_migrations()[:2]:
        conn.executescript(f"BEGIN; {path.read_text(encoding='utf-8')}; INSERT INTO schema_migrations VALUES ({version}, '{name}', 'x'); COMMIT;")
    conn.execute(
        "INSERT INTO content_sets (set_id, kind, source, created_at) VALUES (1, 'models', 'legacy_json:x', 'x')"
    )
    conn.execute("INSERT INTO registry_sets (set_id, default_model_version) VALUES (1, 'Old-1')")
    conn.execute(
        "INSERT INTO models (set_id, model_version, position, display_name, status) VALUES (1, 'Old-1', 0, 'Old 1', 'ACTIVE')"
    )
    conn.execute("INSERT INTO use_case_variants (set_id, model_version, use_case, position) VALUES (1, 'Old-1', 'CDSS', 0)")
    assert db.migrate(conn) == list(range(3, len(db.available_migrations()) + 1))
    registry = repository.load_registry(conn, 1)
    model = registry.models[0]
    assert (model.kind, model.ai_type, model.developer, model.release_date, model.description, model.built_on) == (
        "model", None, None, None, None, (),
    )
    assert model.use_cases[0].function_category is None and model.use_cases[0].declared_contact is None
    with pytest.raises(sqlite3.IntegrityError):
        conn.execute("UPDATE models SET kind = 'service'")
